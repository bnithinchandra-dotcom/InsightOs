import copy
import unittest
from concurrent.futures import ThreadPoolExecutor
from io import BytesIO
from threading import Event, Lock
from types import SimpleNamespace
from unittest.mock import MagicMock, patch

from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.exc import OperationalError, SQLAlchemyError
from sqlalchemy.orm import Session, sessionmaker
from sqlalchemy.pool import StaticPool
from openpyxl import Workbook
import pyarrow as pa
import pyarrow.parquet as pq

from app.database import Base
from app.models import Dataset, DatasetFile, Project, User
from app import main, normalization


class MissingObject(Exception):
    def __init__(self, code: str):
        self.code = code


class FakeStorage:
    def __init__(self):
        self.objects = {}

    def bucket_exists(self, _bucket):
        return any(bucket == _bucket for bucket, _key in self.objects)

    def make_bucket(self, _bucket):
        return None

    def stat_object(self, bucket, key):
        try:
            stored = self.objects[(bucket, key)]
        except KeyError:
            raise MissingObject("NoSuchKey") from None
        return SimpleNamespace(
            size=len(stored["content"]),
            metadata=(
                {"x-amz-meta-sha256": stored["metadata"]["sha256"]}
                if "sha256" in stored["metadata"]
                else {}
            ),
        )

    def get_object(self, bucket, key):
        return BytesIO(self.objects[(bucket, key)]["content"])

    def put_object(
        self,
        bucket,
        key,
        stream,
        _size,
        content_type=None,
        metadata=None,
    ):
        self.objects[(bucket, key)] = {
            "content": stream.read(),
            "content_type": content_type,
            "metadata": metadata or {},
        }

    def remove_object(self, bucket, key):
        self.objects.pop((bucket, key), None)

    def list_objects(self, bucket, prefix, recursive):
        del recursive
        return [
            SimpleNamespace(object_name=key)
            for object_bucket, key in self.objects
            if object_bucket == bucket and key.startswith(prefix)
        ]

    def remove_objects(self, bucket, objects):
        for obj in objects:
            self.remove_object(bucket, obj.name)
        return []


class UploadFlowTests(unittest.TestCase):
    def setUp(self):
        self.engine = create_engine(
            "sqlite://",
            connect_args={"check_same_thread": False},
            poolclass=StaticPool,
        )
        Base.metadata.create_all(self.engine)
        self.session_factory = sessionmaker(bind=self.engine, expire_on_commit=False)
        session = self.session_factory()
        user = User(email="upload-test@example.invalid")
        session.add(user)
        session.flush()
        project = Project(name="Test project", user_id=user.id)
        session.add(project)
        session.flush()
        self.project_id = project.id
        dataset = Dataset(project_id=project.id, name="Test dataset", status="Created")
        session.add(dataset)
        session.commit()
        self.dataset_id = dataset.id
        session.close()

        self.storage = FakeStorage()
        self.real_release_upload_lock = main._release_upload_lock

        def acquire_upload_lock(session, connection, _dataset_id, _key):
            self.assertIs(session.get_bind(), connection)
            return 1

        def release_upload_lock(session, connection, _lock_id):
            self.assertIs(session.get_bind(), connection)

        self.patches = [
            patch.object(main, "SessionLocal", self.session_factory),
            patch.object(
                main,
                "_try_acquire_upload_lock",
                side_effect=acquire_upload_lock,
            ),
            patch.object(
                main,
                "_release_upload_lock",
                side_effect=release_upload_lock,
            ),
            patch.object(main, "get_storage_client", return_value=self.storage),
            patch.object(main, "S3Error", MissingObject),
            patch.object(normalization, "S3Error", MissingObject),
        ]
        for active_patch in self.patches:
            active_patch.start()
        self.client = TestClient(main.app)

    def tearDown(self):
        for active_patch in reversed(self.patches):
            active_patch.stop()
        Base.metadata.drop_all(self.engine)
        self.engine.dispose()

    def test_upload_parses_persists_and_preserves_original_bytes(self):
        content = b'"id",name\r\n1,Ada\r\n2,Lin\r\n'
        response = self.client.post(
            f"/api/v1/datasets/{self.dataset_id}/files",
            files={"file": ("people.csv", content, "text/csv")},
        )
        self.assertEqual(response.status_code, 201, response.text)
        payload = response.json()
        self.assertEqual(payload["dataset_status"], "Uploaded")
        self.assertEqual(payload["file"]["detected_format"], "csv")
        self.assertEqual(payload["file"]["parsing_result"]["row_count"], 2)
        self.assertEqual(payload["file"]["profile_result"]["row_count"], 2)
        self.assertEqual(payload["file"]["status"], "Ready")
        self.assertEqual(len(self.storage.objects), 1)
        stored = next(iter(self.storage.objects.values()))
        self.assertEqual(stored["content"], content)

        session = self.session_factory()
        dataset = session.get(Dataset, self.dataset_id)
        dataset_file = session.query(DatasetFile).one()
        self.assertEqual(dataset.status, "Uploaded")
        self.assertEqual(dataset_file.detected_format, "csv")
        self.assertEqual(dataset_file.parsing_result["columns"][0]["name"], "id")
        self.assertEqual(dataset_file.profile_result["column_count"], 2)
        self.assertEqual(dataset_file.status, "Ready")
        session.close()

        files_response = self.client.get(
            f"/api/v1/datasets/{self.dataset_id}/files"
        )
        self.assertEqual(files_response.status_code, 200)
        self.assertEqual(
            files_response.json()[0]["parsing_result"]["row_count"],
            2,
        )
        self.assertEqual(
            files_response.json()[0]["profile_result"]["row_count"],
            2,
        )
        profile_response = self.client.get(
            f"/api/v1/datasets/{self.dataset_id}/files/{payload['file']['id']}/profile"
        )
        self.assertEqual(profile_response.status_code, 200, profile_response.text)
        self.assertEqual(
            profile_response.json()["profile_result"]["row_count"],
            2,
        )

    def test_parse_failure_sets_failed_without_storing_or_recording_file(self):
        response = self.client.post(
            f"/api/v1/datasets/{self.dataset_id}/files",
            files={"file": ("wrong.csv", b'{"records":[{"x":1}]}', "text/csv")},
        )
        self.assertEqual(response.status_code, 422, response.text)
        self.assertEqual(response.json()["detail"]["code"], "content_mismatch")
        self.assertEqual(self.storage.objects, {})
        session = self.session_factory()
        self.assertEqual(session.get(Dataset, self.dataset_id).status, "Failed")
        dataset_file = session.query(DatasetFile).one()
        self.assertEqual(dataset_file.status, "Failed")
        self.assertEqual(dataset_file.error_code, "content_mismatch")
        file_id = dataset_file.id
        session.close()
        profile_response = self.client.get(
            f"/api/v1/datasets/{self.dataset_id}/files/{file_id}/profile"
        )
        self.assertEqual(profile_response.status_code, 409)
        self.assertEqual(
            profile_response.json()["detail"]["code"],
            "profile_unavailable",
        )

    def test_successful_ingestion_for_all_six_formats_and_replay(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["id", "name"])
        sheet.append([1, "Ada"])
        second_sheet = workbook.create_sheet("guests")
        second_sheet.append(["id", "name"])
        second_sheet.append([2, "Lin"])
        xlsx_buffer = BytesIO()
        workbook.save(xlsx_buffer)

        parquet_buffer = BytesIO()
        pq.write_table(
            pa.table({"id": [1, 2], "name": ["Ada", "Lin"]}),
            parquet_buffer,
        )
        cases = [
            ("people.csv", b"id,name\n1,Ada\n2,Lin\n", "text/csv", "csv"),
            (
                "people.tsv",
                b"id\tname\n1\tAda\n2\tLin\n",
                "text/tab-separated-values",
                "tsv",
            ),
            (
                "people.xlsx",
                xlsx_buffer.getvalue(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                "xlsx",
            ),
            (
                "people.json",
                b'[{"id":1,"name":"Ada"},{"id":2,"name":"Lin"}]',
                "application/json",
                "json",
            ),
            (
                "people.parquet",
                parquet_buffer.getvalue(),
                "application/vnd.apache.parquet",
                "parquet",
            ),
            (
                "people.xml",
                b"<people><person><id>1</id></person><person><id>2</id></person></people>",
                "application/xml",
                "xml",
            ),
        ]
        for filename, content, content_type, expected_format in cases:
            with self.subTest(format=expected_format):
                response = self.client.post(
                    f"/api/v1/datasets/{self.dataset_id}/files",
                    files={"file": (filename, content, content_type)},
                )
                self.assertEqual(response.status_code, 201, response.text)
                payload = response.json()
                self.assertEqual(payload["file"]["detected_format"], expected_format)
                self.assertEqual(payload["file"]["status"], "Ready")
                self.assertEqual(
                    payload["file"]["profile_result"]["detected_format"],
                    expected_format,
                )
                self.assertEqual(
                    payload["file"]["profile_result"]["row_count"],
                    payload["file"]["parsing_result"]["row_count"],
                )
                self.assertGreaterEqual(
                    payload["file"]["parsing_result"]["row_count"],
                    1,
                )
                stored = self.storage.objects[
                    ("insightos-raw", payload["file"]["storage_key"])
                ]
                self.assertEqual(stored["content"], content)

        self.assertEqual(len(self.storage.objects), 6)
        original = cases[0]
        first = self.client.post(
            f"/api/v1/datasets/{self.dataset_id}/files",
            files={"file": original[:3]},
        )
        replay = self.client.post(
            f"/api/v1/datasets/{self.dataset_id}/files",
            files={"file": original[:3]},
        )
        self.assertEqual(first.status_code, 201, first.text)
        self.assertEqual(replay.status_code, 201, replay.text)
        self.assertEqual(first.json()["file"]["id"], replay.json()["file"]["id"])
        session = self.session_factory()
        self.assertEqual(session.query(DatasetFile).count(), 6)
        session.close()

    def test_invalid_profile_metadata_fails_ingestion_without_storing_object(self):
        class InvalidProfileResult:
            profile_result = {"version": 2}

            @staticmethod
            def as_dict():
                return {
                    "detected_format": "csv",
                    "columns": [
                        {
                            "name": "value",
                            "position": 0,
                            "physical_type": "string",
                            "physical_types": ["string"],
                            "missing_values": 0,
                            "empty_values": 0,
                        }
                    ],
                    "row_count": 1,
                    "metadata": {},
                }

        with patch.object(
            main,
            "parse_dataset_file",
            return_value=InvalidProfileResult(),
        ):
            response = self.client.post(
                f"/api/v1/datasets/{self.dataset_id}/files",
                files={"file": ("invalid-profile.csv", b"value\nx\n", "text/csv")},
            )

        self.assertEqual(response.status_code, 500, response.text)
        self.assertEqual(self.storage.objects, {})
        session = self.session_factory()
        dataset_file = session.query(DatasetFile).one()
        self.assertEqual(dataset_file.status, "Failed")
        self.assertEqual(dataset_file.error_code, "invalid_profile_metadata")
        self.assertIsNone(dataset_file.profile_result)
        session.close()

    def test_normalization_api_is_idempotent_and_keeps_ingestion_state_separate(self):
        content = b"code,value\n001,1.25\n002,2.50\n"
        upload = self.client.post(
            f"/api/v1/datasets/{self.dataset_id}/files",
            files={"file": ("numbers.csv", content, "text/csv")},
        )
        self.assertEqual(upload.status_code, 201, upload.text)
        file_id = upload.json()["file"]["id"]

        first = self.client.post(
            f"/api/v1/datasets/{self.dataset_id}/files/{file_id}/normalize"
        )
        self.assertEqual(first.status_code, 200, first.text)
        first_payload = first.json()
        self.assertEqual(first_payload["status"], "Ready")
        self.assertEqual(first_payload["normalization_result"]["row_count"], 2)
        self.assertEqual(
            first_payload["normalization_result"]["tables"][0]["schema"][0],
            {"name": "code", "type": "string"},
        )

        stored_original = self.storage.objects[
            ("insightos-raw", upload.json()["file"]["storage_key"])
        ]["content"]
        self.assertEqual(stored_original, content)
        normalized_count = len(
            [
                key
                for key in self.storage.objects
                if key[0] == "insightos-processed"
            ]
        )
        second = self.client.post(
            f"/api/v1/datasets/{self.dataset_id}/files/{file_id}/normalize"
        )
        self.assertEqual(second.status_code, 200, second.text)
        self.assertEqual(
            second.json()["normalization_result"]["tables"][0]["object_key"],
            first_payload["normalization_result"]["tables"][0]["object_key"],
        )
        self.assertEqual(
            len(
                [
                    key
                    for key in self.storage.objects
                    if key[0] == "insightos-processed"
                ]
            ),
            normalized_count,
        )

        status = self.client.get(
            f"/api/v1/datasets/{self.dataset_id}/files/{file_id}/normalization"
        )
        self.assertEqual(status.status_code, 200, status.text)
        self.assertEqual(status.json()["status"], "Ready")
        session = self.session_factory()
        dataset_file = session.get(DatasetFile, file_id)
        self.assertEqual(dataset_file.status, "Ready")
        self.assertEqual(dataset_file.normalization_status, "Ready")
        session.close()

    def test_quality_analysis_persists_contract_without_changing_existing_apis_or_data(self):
        content = b"code,value\n001,1.25\n002,2.50\n"
        upload = self.client.post(
            f"/api/v1/datasets/{self.dataset_id}/files",
            files={"file": ("quality.csv", content, "text/csv")},
        )
        self.assertEqual(upload.status_code, 201, upload.text)
        file_id = upload.json()["file"]["id"]
        normalized = self.client.post(
            f"/api/v1/datasets/{self.dataset_id}/files/{file_id}/normalize"
        )
        self.assertEqual(normalized.status_code, 200, normalized.text)
        normalized_body = normalized.json()
        self.assertEqual(
            set(normalized_body),
            {
                "dataset_id",
                "file_id",
                "status",
                "normalization_result",
                "error",
            },
        )

        profile = self.client.get(
            f"/api/v1/datasets/{self.dataset_id}/files/{file_id}/profile"
        )
        self.assertEqual(profile.status_code, 200, profile.text)
        self.assertEqual(
            set(profile.json()),
            {"dataset_id", "file_id", "profile_result"},
        )
        profile_before = profile.json()
        storage_before = copy.deepcopy(self.storage.objects)

        first = self.client.post(
            f"/api/v1/datasets/{self.dataset_id}/files/{file_id}/quality-analysis"
        )
        self.assertEqual(first.status_code, 200, first.text)
        first_body = first.json()
        self.assertEqual(first_body["status"], "PartiallyCompleted")
        self.assertTrue(first_body["report_is_current"])
        self.assertEqual(first_body["quality_report"]["schema_version"], 1)
        self.assertEqual(first_body["quality_report"]["findings"], [])
        self.assertEqual(
            first_body["quality_report"]["coverage"]["checks"][0]["status"],
            "Skipped",
        )

        repeated = self.client.post(
            f"/api/v1/datasets/{self.dataset_id}/files/{file_id}/quality-analysis"
        )
        self.assertEqual(repeated.status_code, 200, repeated.text)
        self.assertEqual(
            repeated.json()["quality_report"],
            first_body["quality_report"],
        )
        status = self.client.get(
            f"/api/v1/datasets/{self.dataset_id}/files/{file_id}/quality-analysis"
        )
        self.assertEqual(status.status_code, 200, status.text)
        self.assertEqual(status.json(), repeated.json())
        self.assertEqual(self.storage.objects, storage_before)
        self.assertEqual(
            self.client.get(
                f"/api/v1/datasets/{self.dataset_id}/files/{file_id}/profile"
            ).json(),
            profile_before,
        )
        normalization_status = self.client.get(
            f"/api/v1/datasets/{self.dataset_id}/files/{file_id}/normalization"
        )
        self.assertEqual(
            normalization_status.json()["normalization_result"],
            normalized_body["normalization_result"],
        )

        session = self.session_factory()
        dataset_file = session.get(DatasetFile, file_id)
        self.assertEqual(dataset_file.status, "Ready")
        self.assertEqual(dataset_file.normalization_status, "Ready")
        self.assertEqual(
            dataset_file.quality_analysis_status,
            "PartiallyCompleted",
        )
        self.assertEqual(
            dataset_file.quality_analysis_key,
            first_body["analysis_key"],
        )
        session.close()

    def test_quality_analysis_failure_keeps_previous_report_and_safe_error(self):
        upload = self.client.post(
            f"/api/v1/datasets/{self.dataset_id}/files",
            files={"file": ("quality-failure.csv", b"value\n1\n", "text/csv")},
        )
        self.assertEqual(upload.status_code, 201, upload.text)
        file_id = upload.json()["file"]["id"]
        normalized = self.client.post(
            f"/api/v1/datasets/{self.dataset_id}/files/{file_id}/normalize"
        )
        self.assertEqual(normalized.status_code, 200, normalized.text)
        successful = self.client.post(
            f"/api/v1/datasets/{self.dataset_id}/files/{file_id}/quality-analysis"
        )
        self.assertEqual(successful.status_code, 200, successful.text)
        previous_report = successful.json()["quality_report"]

        with patch.object(
            main,
            "verify_normalization_result",
            side_effect=OSError("private source value"),
        ):
            failed = self.client.post(
                f"/api/v1/datasets/{self.dataset_id}/files/{file_id}/quality-analysis",
                json={"recompute": True},
            )
        self.assertEqual(failed.status_code, 502, failed.text)
        self.assertNotIn("private source value", failed.text)

        status = self.client.get(
            f"/api/v1/datasets/{self.dataset_id}/files/{file_id}/quality-analysis"
        )
        self.assertEqual(status.status_code, 200, status.text)
        self.assertEqual(status.json()["status"], "Failed")
        self.assertTrue(status.json()["report_is_current"])
        self.assertEqual(status.json()["quality_report"], previous_report)
        self.assertEqual(status.json()["error"]["code"], "storage_failure")
        self.assertNotIn("private source value", status.text)

    def test_unexpected_quality_analysis_failure_keeps_previous_report(self):
        upload = self.client.post(
            f"/api/v1/datasets/{self.dataset_id}/files",
            files={
                "file": (
                    "quality-unexpected-failure.csv",
                    b"value\n1\n",
                    "text/csv",
                )
            },
        )
        self.assertEqual(upload.status_code, 201, upload.text)
        file_id = upload.json()["file"]["id"]
        normalized = self.client.post(
            f"/api/v1/datasets/{self.dataset_id}/files/{file_id}/normalize"
        )
        self.assertEqual(normalized.status_code, 200, normalized.text)
        successful = self.client.post(
            f"/api/v1/datasets/{self.dataset_id}/files/{file_id}/quality-analysis"
        )
        self.assertEqual(successful.status_code, 200, successful.text)
        previous_report = successful.json()["quality_report"]

        with patch.object(
            main,
            "build_contract_report",
            side_effect=RuntimeError("private cell value"),
        ):
            failed = self.client.post(
                f"/api/v1/datasets/{self.dataset_id}/files/{file_id}/quality-analysis",
                json={"recompute": True},
            )
        self.assertEqual(failed.status_code, 500, failed.text)
        self.assertEqual(
            failed.json()["detail"]["code"],
            "quality_analysis_failed",
        )
        self.assertNotIn("private cell value", failed.text)

        state = self.client.get(
            f"/api/v1/datasets/{self.dataset_id}/files/{file_id}/quality-analysis"
        )
        self.assertEqual(state.status_code, 200, state.text)
        self.assertEqual(state.json()["status"], "Failed")
        self.assertTrue(state.json()["report_is_current"])
        self.assertEqual(state.json()["quality_report"], previous_report)
        self.assertEqual(
            state.json()["quality_report"]["analysis"]["analysis_key"],
            successful.json()["analysis_key"],
        )
        self.assertNotIn("private cell value", state.text)

    def test_quality_report_remains_available_and_stale_after_renormalization(self):
        upload = self.client.post(
            f"/api/v1/datasets/{self.dataset_id}/files",
            files={
                "file": (
                    "quality-renormalize.csv",
                    b"value\n1\n2\n3\n",
                    "text/csv",
                )
            },
        )
        self.assertEqual(upload.status_code, 201, upload.text)
        file_id = upload.json()["file"]["id"]
        first_normalization = self.client.post(
            f"/api/v1/datasets/{self.dataset_id}/files/{file_id}/normalize"
        )
        self.assertEqual(first_normalization.status_code, 200)
        first_analysis = self.client.post(
            f"/api/v1/datasets/{self.dataset_id}/files/{file_id}/quality-analysis"
        )
        self.assertEqual(first_analysis.status_code, 200, first_analysis.text)
        previous_report = first_analysis.json()["quality_report"]
        previous_key = previous_report["analysis"]["analysis_key"]

        with patch.dict("os.environ", {"NORMALIZATION_COMPRESSION": "snappy"}):
            normalized_again = self.client.post(
                f"/api/v1/datasets/{self.dataset_id}/files/{file_id}/normalize"
            )
        self.assertEqual(normalized_again.status_code, 200, normalized_again.text)
        self.assertNotEqual(
            normalized_again.json()["normalization_result"]["configuration"][
                "configuration_sha256"
            ],
            first_normalization.json()["normalization_result"]["configuration"][
                "configuration_sha256"
            ],
        )

        stale = self.client.get(
            f"/api/v1/datasets/{self.dataset_id}/files/{file_id}/quality-analysis"
        )
        self.assertEqual(stale.status_code, 200, stale.text)
        self.assertFalse(stale.json()["report_is_current"])
        self.assertEqual(stale.json()["quality_report"], previous_report)
        self.assertEqual(
            stale.json()["quality_report"]["analysis"]["analysis_key"],
            previous_key,
        )

        new_analysis = self.client.post(
            f"/api/v1/datasets/{self.dataset_id}/files/{file_id}/quality-analysis"
        )
        self.assertEqual(new_analysis.status_code, 200, new_analysis.text)
        self.assertTrue(new_analysis.json()["report_is_current"])
        self.assertNotEqual(
            new_analysis.json()["quality_report"]["analysis"]["analysis_key"],
            previous_key,
        )

    def test_quality_report_survives_failed_normalization_retry(self):
        upload = self.client.post(
            f"/api/v1/datasets/{self.dataset_id}/files",
            files={
                "file": (
                    "quality-failed-normalize-retry.csv",
                    b"value\n1\n",
                    "text/csv",
                )
            },
        )
        self.assertEqual(upload.status_code, 201, upload.text)
        file_id = upload.json()["file"]["id"]
        normalized = self.client.post(
            f"/api/v1/datasets/{self.dataset_id}/files/{file_id}/normalize"
        )
        self.assertEqual(normalized.status_code, 200, normalized.text)
        analysis = self.client.post(
            f"/api/v1/datasets/{self.dataset_id}/files/{file_id}/quality-analysis"
        )
        self.assertEqual(analysis.status_code, 200, analysis.text)
        previous_report = analysis.json()["quality_report"]
        output_key = normalized.json()["normalization_result"]["tables"][0][
            "object_key"
        ]
        self.storage.remove_object("insightos-processed", output_key)

        original_put = self.storage.put_object

        def fail_processed(bucket, *args, **kwargs):
            if bucket == "insightos-processed":
                raise OSError("injected retry failure")
            return original_put(bucket, *args, **kwargs)

        with patch.object(self.storage, "put_object", side_effect=fail_processed):
            failed = self.client.post(
                f"/api/v1/datasets/{self.dataset_id}/files/{file_id}/normalize"
            )
        self.assertEqual(failed.status_code, 502, failed.text)

        stale = self.client.get(
            f"/api/v1/datasets/{self.dataset_id}/files/{file_id}/quality-analysis"
        )
        self.assertEqual(stale.status_code, 200, stale.text)
        self.assertFalse(stale.json()["report_is_current"])
        self.assertEqual(stale.json()["quality_report"], previous_report)

        retried = self.client.post(
            f"/api/v1/datasets/{self.dataset_id}/files/{file_id}/normalize"
        )
        self.assertEqual(retried.status_code, 200, retried.text)
        current = self.client.get(
            f"/api/v1/datasets/{self.dataset_id}/files/{file_id}/quality-analysis"
        )
        self.assertEqual(current.status_code, 200, current.text)
        self.assertTrue(current.json()["report_is_current"])
        self.assertEqual(current.json()["quality_report"], previous_report)

    def test_quality_analysis_persists_running_and_rejects_duplicate_request(self):
        upload = self.client.post(
            f"/api/v1/datasets/{self.dataset_id}/files",
            files={"file": ("quality-concurrent.csv", b"value\n1\n", "text/csv")},
        )
        self.assertEqual(upload.status_code, 201, upload.text)
        file_id = upload.json()["file"]["id"]
        normalized = self.client.post(
            f"/api/v1/datasets/{self.dataset_id}/files/{file_id}/normalize"
        )
        self.assertEqual(normalized.status_code, 200, normalized.text)

        held = set()
        lock_ids = {}
        next_lock_id = 800
        lock_guard = Lock()
        running = Event()
        allow_analysis_to_finish = Event()
        original_build = main.build_contract_report

        def acquire(session, connection, dataset_id, key):
            nonlocal next_lock_id
            self.assertIs(session.get_bind(), connection)
            self.assertEqual(dataset_id, self.dataset_id)
            self.assertEqual(key, f"quality-analysis:{file_id}")
            with lock_guard:
                if key in held:
                    return None
                held.add(key)
                next_lock_id += 1
                lock_ids[next_lock_id] = key
                return next_lock_id

        def release(session, connection, lock_id):
            self.assertIs(session.get_bind(), connection)
            with lock_guard:
                held.discard(lock_ids.pop(lock_id))

        def pause_report(*args, **kwargs):
            running.set()
            if not allow_analysis_to_finish.wait(timeout=10):
                raise TimeoutError("test did not release quality analysis")
            return original_build(*args, **kwargs)

        with (
            patch.object(main, "_try_acquire_upload_lock", side_effect=acquire),
            patch.object(main, "_release_upload_lock", side_effect=release),
            patch.object(main, "build_contract_report", side_effect=pause_report),
            ThreadPoolExecutor(max_workers=1) as executor,
        ):
            future = executor.submit(
                self.client.post,
                f"/api/v1/datasets/{self.dataset_id}/files/{file_id}/quality-analysis",
            )
            try:
                self.assertTrue(running.wait(timeout=5))
                state = self.client.get(
                    f"/api/v1/datasets/{self.dataset_id}/files/{file_id}/quality-analysis"
                )
                self.assertEqual(state.status_code, 200, state.text)
                self.assertEqual(state.json()["status"], "Running")
                duplicate = self.client.post(
                    f"/api/v1/datasets/{self.dataset_id}/files/{file_id}/quality-analysis"
                )
                self.assertEqual(duplicate.status_code, 409, duplicate.text)
                self.assertEqual(
                    duplicate.json()["detail"]["code"],
                    "quality_analysis_in_progress",
                )
            finally:
                allow_analysis_to_finish.set()
            result = future.result(timeout=10)
        self.assertEqual(result.status_code, 200, result.text)
        self.assertEqual(result.json()["status"], "PartiallyCompleted")

    def test_quality_analysis_rejects_unready_normalization_without_mutation(self):
        upload = self.client.post(
            f"/api/v1/datasets/{self.dataset_id}/files",
            files={"file": ("not-normalized.csv", b"value\n1\n", "text/csv")},
        )
        self.assertEqual(upload.status_code, 201, upload.text)
        file_id = upload.json()["file"]["id"]
        response = self.client.post(
            f"/api/v1/datasets/{self.dataset_id}/files/{file_id}/quality-analysis"
        )
        self.assertEqual(response.status_code, 409, response.text)
        self.assertEqual(
            response.json()["detail"]["code"],
            "normalization_unavailable",
        )
        session = self.session_factory()
        dataset_file = session.get(DatasetFile, file_id)
        self.assertEqual(dataset_file.quality_analysis_status, "NotStarted")
        self.assertIsNone(dataset_file.quality_report)
        session.close()

    def test_quality_analysis_conflicts_with_dataset_deletion_while_running(self):
        upload = self.client.post(
            f"/api/v1/datasets/{self.dataset_id}/files",
            files={"file": ("quality-delete.csv", b"value\n1\n", "text/csv")},
        )
        self.assertEqual(upload.status_code, 201, upload.text)
        file_id = upload.json()["file"]["id"]
        normalized = self.client.post(
            f"/api/v1/datasets/{self.dataset_id}/files/{file_id}/normalize"
        )
        self.assertEqual(normalized.status_code, 200, normalized.text)

        held = set()
        lock_ids = {}
        analyzed = Event()
        allow_analysis_to_finish = Event()
        original_verify = main.verify_normalization_result

        def acquire(session, connection, _dataset_id, key):
            self.assertIs(session.get_bind(), connection)
            self.assertEqual(_dataset_id, self.dataset_id)
            if key in held:
                return None
            held.add(key)
            lock_id = len(held) + 700
            lock_ids[lock_id] = key
            return lock_id

        def release(session, connection, lock_id):
            self.assertIs(session.get_bind(), connection)
            held.discard(lock_ids.pop(lock_id))

        def pause_verification(storage, bucket, result):
            analyzed.set()
            if not allow_analysis_to_finish.wait(timeout=10):
                raise TimeoutError("test did not release quality analysis")
            return original_verify(storage, bucket, result)

        with (
            patch.object(main, "_try_acquire_upload_lock", side_effect=acquire),
            patch.object(main, "_release_upload_lock", side_effect=release),
            patch.object(
                main,
                "verify_normalization_result",
                side_effect=pause_verification,
            ),
            ThreadPoolExecutor(max_workers=1) as executor,
        ):
            analysis_future = executor.submit(
                self.client.post,
                f"/api/v1/datasets/{self.dataset_id}/files/{file_id}/quality-analysis",
            )
            try:
                self.assertTrue(analyzed.wait(timeout=5))
                deleted = self.client.delete(
                    f"/api/v1/datasets/{self.dataset_id}"
                )
                self.assertEqual(deleted.status_code, 409, deleted.text)
                self.assertEqual(
                    deleted.json()["detail"]["code"],
                    "quality_analysis_in_progress",
                )
            finally:
                allow_analysis_to_finish.set()
            analysis = analysis_future.result(timeout=10)
        self.assertEqual(analysis.status_code, 200, analysis.text)

    def test_normalization_storage_failure_does_not_fail_ingestion_and_can_retry(self):
        content = b"value\n1\n"
        upload = self.client.post(
            f"/api/v1/datasets/{self.dataset_id}/files",
            files={"file": ("retry.csv", content, "text/csv")},
        )
        self.assertEqual(upload.status_code, 201, upload.text)
        file_id = upload.json()["file"]["id"]
        original_put = self.storage.put_object

        def fail_processed(bucket, *args, **kwargs):
            if bucket == "insightos-processed":
                raise OSError("injected normalization storage failure")
            return original_put(bucket, *args, **kwargs)

        with patch.object(self.storage, "put_object", side_effect=fail_processed):
            failed = self.client.post(
                f"/api/v1/datasets/{self.dataset_id}/files/{file_id}/normalize"
            )
        self.assertEqual(failed.status_code, 502, failed.text)
        session = self.session_factory()
        dataset_file = session.get(DatasetFile, file_id)
        self.assertEqual(dataset_file.status, "Ready")
        self.assertEqual(dataset_file.normalization_status, "Failed")
        self.assertEqual(dataset_file.normalization_error_code, "storage_failure")
        session.close()

        retried = self.client.post(
            f"/api/v1/datasets/{self.dataset_id}/files/{file_id}/normalize"
        )
        self.assertEqual(retried.status_code, 200, retried.text)
        self.assertEqual(retried.json()["status"], "Ready")

    def test_failed_normalization_retry_does_not_return_stale_success_result(self):
        upload = self.client.post(
            f"/api/v1/datasets/{self.dataset_id}/files",
            files={"file": ("retry-state.csv", b"value\n1\n", "text/csv")},
        )
        self.assertEqual(upload.status_code, 201, upload.text)
        file_id = upload.json()["file"]["id"]
        successful = self.client.post(
            f"/api/v1/datasets/{self.dataset_id}/files/{file_id}/normalize"
        )
        self.assertEqual(successful.status_code, 200, successful.text)
        self.assertIsNotNone(successful.json()["normalization_result"])
        output_key = successful.json()["normalization_result"]["tables"][0]["object_key"]
        self.storage.remove_object("insightos-processed", output_key)

        original_put = self.storage.put_object

        def fail_processed(bucket, *args, **kwargs):
            if bucket == "insightos-processed":
                raise OSError("injected retry failure")
            return original_put(bucket, *args, **kwargs)

        with patch.object(self.storage, "put_object", side_effect=fail_processed):
            failed = self.client.post(
                f"/api/v1/datasets/{self.dataset_id}/files/{file_id}/normalize"
            )
        self.assertEqual(failed.status_code, 502, failed.text)
        self.assertEqual(failed.json()["detail"]["code"], "storage_failure")
        status = self.client.get(
            f"/api/v1/datasets/{self.dataset_id}/files/{file_id}/normalization"
        )
        self.assertEqual(status.json()["status"], "Failed")
        self.assertIsNone(status.json()["normalization_result"])

    def test_concurrent_normalization_returns_in_progress_without_removing_output(self):
        upload = self.client.post(
            f"/api/v1/datasets/{self.dataset_id}/files",
            files={"file": ("concurrent.csv", b"value\n1\n", "text/csv")},
        )
        self.assertEqual(upload.status_code, 201, upload.text)
        file_id = upload.json()["file"]["id"]
        lock_guard = Lock()
        held = set()
        lock_keys = {}
        output_written = Event()
        allow_first_to_finish = Event()
        original_put = self.storage.put_object

        def try_lock(session, connection, _dataset_id, key):
            self.assertIs(session.get_bind(), connection)
            self.assertIn(
                key,
                {
                    f"normalization:{file_id}",
                    f"quality-analysis:{file_id}",
                },
            )
            with lock_guard:
                if key in held:
                    return None
                held.add(key)
                lock_id = 303 if key.startswith("normalization:") else 304
                lock_keys[lock_id] = key
                return lock_id

        def release_lock(session, connection, lock_id):
            self.assertIs(session.get_bind(), connection)
            self.assertIn(lock_id, {303, 304})
            with lock_guard:
                held.discard(lock_keys.pop(lock_id))

        def put_and_pause(bucket, *args, **kwargs):
            original_put(bucket, *args, **kwargs)
            if bucket == "insightos-processed":
                output_written.set()
                if not allow_first_to_finish.wait(timeout=10):
                    raise TimeoutError("test did not release normalization")

        with (
            patch.object(main, "_try_acquire_upload_lock", side_effect=try_lock),
            patch.object(main, "_release_upload_lock", side_effect=release_lock),
            patch.object(self.storage, "put_object", side_effect=put_and_pause),
            ThreadPoolExecutor(max_workers=1) as executor,
        ):
            first_future = executor.submit(
                self.client.post,
                f"/api/v1/datasets/{self.dataset_id}/files/{file_id}/normalize",
            )
            try:
                self.assertTrue(output_written.wait(timeout=5))
                second = self.client.post(
                    f"/api/v1/datasets/{self.dataset_id}/files/{file_id}/normalize"
                )
                self.assertEqual(second.status_code, 409, second.text)
                self.assertEqual(
                    second.json()["detail"]["code"],
                    "normalization_in_progress",
                )
                processed_objects = {
                    key: value
                    for key, value in self.storage.objects.items()
                    if key[0] == "insightos-processed"
                }
                self.assertEqual(len(processed_objects), 1)
            finally:
                allow_first_to_finish.set()
            first = first_future.result(timeout=10)

        self.assertEqual(first.status_code, 200, first.text)
        self.assertEqual(
            {
                key: value
                for key, value in self.storage.objects.items()
                if key[0] == "insightos-processed"
            },
            processed_objects,
        )

    def test_profile_columns_must_match_parsing_metadata(self):
        content = b"first,second\n1,2\n"
        parsed = main.parse_dataset_file(
            "metadata.csv",
            BytesIO(content),
            len(content),
        )
        parsing_result = main.validate_parsing_result(
            "metadata.csv",
            parsed,
            len(content),
        )

        mutations = (
            lambda profile: profile["tables"][0].update(name="unexpected"),
            lambda profile: profile["tables"][0]["columns"][0].update(
                name="different"
            ),
            lambda profile: profile["tables"][0]["columns"][0].update(
                position=1
            ),
            lambda profile: profile["tables"][0]["columns"][0].update(
                physical_types=["integer"]
            ),
        )
        for mutate in mutations:
            with self.subTest(mutate=mutate):
                profile = copy.deepcopy(parsed.profile_result)
                mutate(profile)
                with self.assertRaises(main.InvalidProfileResultError):
                    main.validate_profile_result(profile, parsing_result)

    def test_profile_worksheet_rows_must_match_parsing_metadata(self):
        workbook = Workbook()
        workbook.active.title = "First"
        workbook.active.append(["value"])
        workbook.active.append([1])
        second = workbook.create_sheet("Second")
        second.append(["value"])
        second.append([2])
        content_stream = BytesIO()
        workbook.save(content_stream)
        content = content_stream.getvalue()

        parsed = main.parse_dataset_file(
            "worksheets.xlsx",
            BytesIO(content),
            len(content),
        )
        parsing_result = main.validate_parsing_result(
            "worksheets.xlsx",
            parsed,
            len(content),
        )
        profile = copy.deepcopy(parsed.profile_result)
        profile["tables"][0]["row_count"] = 2
        profile["tables"][1]["row_count"] = 0

        with self.assertRaises(main.InvalidProfileResultError):
            main.validate_profile_result(profile, parsing_result)

    def test_explicit_idempotency_key_rejects_different_content(self):
        headers = {"Idempotency-Key": "client-request-1"}
        first = self.client.post(
            f"/api/v1/datasets/{self.dataset_id}/files",
            files={"file": ("first.csv", b"value\n1\n", "text/csv")},
            headers=headers,
        )
        conflict = self.client.post(
            f"/api/v1/datasets/{self.dataset_id}/files",
            files={"file": ("first.csv", b"value\n2\n", "text/csv")},
            headers=headers,
        )
        self.assertEqual(first.status_code, 201, first.text)
        self.assertEqual(conflict.status_code, 409)
        session = self.session_factory()
        self.assertEqual(session.query(DatasetFile).count(), 1)
        session.close()

    def test_concurrent_same_key_request_cannot_clean_up_first_request_object(self):
        key = "concurrent-upload-1"
        content = b"value\n1\n"
        lock_guard = Lock()
        first_request_holds_lock = False
        object_written = Event()
        allow_first_request_to_continue = Event()
        remove_calls = []
        original_put = self.storage.put_object
        original_remove = self.storage.remove_object

        def try_lock(session, connection, _dataset_id, _key):
            self.assertIs(session.get_bind(), connection)
            nonlocal first_request_holds_lock
            with lock_guard:
                if first_request_holds_lock:
                    return None
                first_request_holds_lock = True
                return 101

        def release_lock(session, connection, _lock_id):
            self.assertIs(session.get_bind(), connection)
            nonlocal first_request_holds_lock
            with lock_guard:
                first_request_holds_lock = False

        def put_then_pause(*args, **kwargs):
            original_put(*args, **kwargs)
            object_written.set()
            if not allow_first_request_to_continue.wait(timeout=10):
                raise TimeoutError("test did not release the first upload")

        def track_remove(*args, **kwargs):
            remove_calls.append(args[1])
            return original_remove(*args, **kwargs)

        with (
            patch.object(main, "_try_acquire_upload_lock", side_effect=try_lock),
            patch.object(main, "_release_upload_lock", side_effect=release_lock),
            patch.object(self.storage, "put_object", side_effect=put_then_pause),
            patch.object(self.storage, "remove_object", side_effect=track_remove),
            ThreadPoolExecutor(max_workers=1) as executor,
        ):
            first_future = executor.submit(
                self.client.post,
                f"/api/v1/datasets/{self.dataset_id}/files",
                files={"file": ("concurrent.csv", content, "text/csv")},
                headers={"Idempotency-Key": key},
            )
            try:
                self.assertTrue(object_written.wait(timeout=5))
                second = self.client.post(
                    f"/api/v1/datasets/{self.dataset_id}/files",
                    files={"file": ("concurrent.csv", content, "text/csv")},
                    headers={"Idempotency-Key": key},
                )
                self.assertEqual(second.status_code, 409, second.text)
                self.assertEqual(second.json()["detail"]["code"], "upload_in_progress")
                self.assertEqual(remove_calls, [])
                self.assertEqual(len(self.storage.objects), 1)
            finally:
                allow_first_request_to_continue.set()

            first = first_future.result(timeout=10)

        self.assertEqual(first.status_code, 201, first.text)
        self.assertEqual(remove_calls, [])
        self.assertEqual(len(self.storage.objects), 1)

    def test_pinned_connection_identity_survives_commits_and_release(self):
        class ConnectionCheckingSession(Session):
            def commit(inner_self):
                expected_connection = getattr(
                    inner_self,
                    "expected_upload_connection",
                    None,
                )
                if expected_connection is not None:
                    self.assertIs(inner_self.get_bind(), expected_connection)
                super().commit()

        checking_factory = sessionmaker(
            bind=self.engine,
            class_=ConnectionCheckingSession,
            expire_on_commit=False,
        )
        captured = {}

        def acquire(session, connection, _dataset_id, _key):
            captured["connection"] = connection
            session.expected_upload_connection = connection
            self.assertIs(session.get_bind(), connection)
            return 202

        def release(session, connection, lock_id):
            self.assertEqual(lock_id, 202)
            self.assertIs(connection, captured["connection"])
            self.assertIs(session.get_bind(), connection)

        with (
            patch.object(main, "SessionLocal", checking_factory),
            patch.object(main, "_try_acquire_upload_lock", side_effect=acquire),
            patch.object(main, "_release_upload_lock", side_effect=release),
        ):
            response = self.client.post(
                f"/api/v1/datasets/{self.dataset_id}/files",
                files={"file": ("pinned.csv", b"value\n1\n", "text/csv")},
                headers={"Idempotency-Key": "pinned-connection"},
            )

        self.assertEqual(response.status_code, 201, response.text)
        self.assertIsNotNone(captured["connection"])

    def test_failed_unlock_invalidates_connection_and_raises(self):
        session = MagicMock()
        connection = MagicMock()
        session.get_bind.return_value = connection
        session.scalar.return_value = False

        with self.assertRaisesRegex(RuntimeError, "connection was discarded"):
            self.real_release_upload_lock(session, connection, 303)

        connection.invalidate.assert_called_once()
        session.commit.assert_not_called()

    def test_failed_invalidate_detach_and_physical_close_quarantines_connection(self):
        session = MagicMock()
        bind = MagicMock()
        connection = MagicMock()
        proxied_connection = connection.connection
        proxied_connection.is_valid = True
        proxied_connection.is_detached = False
        proxied_connection.dbapi_connection.close.side_effect = OSError(
            "physical close failed"
        )
        connection.invalidated = False
        connection.invalidate.side_effect = SQLAlchemyError("invalidation failed")
        connection.detach.side_effect = SQLAlchemyError("detachment failed")
        session.get_bind.return_value = bind
        bind.connect.return_value = connection
        cleanup_error = RuntimeError("original upload cleanup failure")

        with patch.object(main, "SessionLocal", return_value=session):
            with self.assertRaises(
                main._UploadConnectionDisposalError
            ) as raised_error:
                with main._pinned_upload_session() as (_, pinned_connection):
                    self.assertIs(pinned_connection, connection)
                    main._discard_upload_connection(connection, cleanup_error)

        disposal_error = raised_error.exception
        self.assertIs(disposal_error.__cause__, cleanup_error)
        self.assertIs(disposal_error.cleanup_error, cleanup_error)
        self.assertIsInstance(disposal_error.invalidate_error, SQLAlchemyError)
        self.assertIsInstance(disposal_error.detach_error, SQLAlchemyError)
        self.assertIsInstance(disposal_error.physical_close_error, OSError)
        connection.close.assert_not_called()
        session.close.assert_called_once()
        self.assertTrue(
            any(
                quarantined is connection
                for quarantined in main._UPLOAD_CONNECTION_QUARANTINE
            )
        )

        with main._UPLOAD_CONNECTION_QUARANTINE_LOCK:
            main._UPLOAD_CONNECTION_QUARANTINE[:] = [
                quarantined
                for quarantined in main._UPLOAD_CONNECTION_QUARANTINE
                if quarantined is not connection
            ]

    def test_failed_invalidate_and_detach_closes_dbapi_before_connection_cleanup(self):
        session = MagicMock()
        bind = MagicMock()
        connection = MagicMock()
        proxied_connection = connection.connection
        proxied_connection.is_valid = True
        proxied_connection.is_detached = False
        connection.invalidated = False
        connection.invalidate.side_effect = SQLAlchemyError("invalidation failed")
        connection.detach.side_effect = SQLAlchemyError("detachment failed")
        session.get_bind.side_effect = [bind, connection]
        session.scalar.return_value = False
        bind.connect.return_value = connection

        with patch.object(main, "SessionLocal", return_value=session):
            with self.assertRaisesRegex(RuntimeError, "connection was discarded"):
                with main._pinned_upload_session() as (_, pinned_connection):
                    self.real_release_upload_lock(
                        session,
                        pinned_connection,
                        303,
                    )

        proxied_connection.dbapi_connection.close.assert_called_once()
        connection.close.assert_called_once()

    def test_invalid_parser_metadata_cannot_mark_ingestion_successful(self):
        class InvalidResult:
            def as_dict(self):
                return {
                    "detected_format": "csv",
                    "columns": [],
                    "row_count": "one",
                    "metadata": {},
                }

        with patch.object(main, "parse_dataset_file", return_value=InvalidResult()):
            response = self.client.post(
                f"/api/v1/datasets/{self.dataset_id}/files",
                files={"file": ("invalid-metadata.csv", b"value\n1\n", "text/csv")},
            )
        self.assertEqual(response.status_code, 500)
        self.assertEqual(self.storage.objects, {})
        session = self.session_factory()
        dataset_file = session.query(DatasetFile).one()
        self.assertEqual(dataset_file.status, "Failed")
        self.assertEqual(dataset_file.error_code, "invalid_parsing_metadata")
        self.assertEqual(session.get(Dataset, self.dataset_id).status, "Failed")
        session.close()

    def test_partial_storage_failure_is_retryable_without_duplicate_objects(self):
        original_put = self.storage.put_object
        should_fail = True

        def put_then_fail(*args, **kwargs):
            nonlocal should_fail
            original_put(*args, **kwargs)
            if should_fail:
                should_fail = False
                raise OSError("injected partial storage failure")

        with patch.object(self.storage, "put_object", side_effect=put_then_fail):
            failed = self.client.post(
                f"/api/v1/datasets/{self.dataset_id}/files",
                files={"file": ("retry.csv", b"value\n1\n", "text/csv")},
            )
        self.assertEqual(failed.status_code, 502)
        session = self.session_factory()
        dataset_file = session.query(DatasetFile).one()
        self.assertEqual(dataset_file.status, "Failed")
        self.assertEqual(session.get(Dataset, self.dataset_id).status, "Failed")
        file_id = dataset_file.id
        session.close()
        self.assertEqual(self.storage.objects, {})

        retried = self.client.post(
            f"/api/v1/datasets/{self.dataset_id}/files",
            files={"file": ("retry.csv", b"value\n1\n", "text/csv")},
        )
        self.assertEqual(retried.status_code, 201, retried.text)
        self.assertEqual(retried.json()["file"]["id"], file_id)
        self.assertEqual(len(self.storage.objects), 1)

    def test_unconfirmed_storage_cleanup_is_recovered_by_same_retry(self):
        original_put = self.storage.put_object
        should_fail = True

        def put_then_fail(*args, **kwargs):
            nonlocal should_fail
            original_put(*args, **kwargs)
            if should_fail:
                should_fail = False
                raise OSError("injected partial storage failure")

        with (
            patch.object(self.storage, "put_object", side_effect=put_then_fail),
            patch.object(
                self.storage,
                "remove_object",
                side_effect=OSError("injected cleanup failure"),
            ),
        ):
            failed = self.client.post(
                f"/api/v1/datasets/{self.dataset_id}/files",
                files={"file": ("recover.csv", b"value\n1\n", "text/csv")},
            )
        self.assertEqual(failed.status_code, 502)
        self.assertEqual(
            failed.json()["detail"]["code"],
            "storage_cleanup_unconfirmed",
        )
        session = self.session_factory()
        dataset_file = session.query(DatasetFile).one()
        self.assertEqual(dataset_file.status, "Processing")
        self.assertEqual(session.get(Dataset, self.dataset_id).status, "Processing")
        session.close()
        self.assertEqual(len(self.storage.objects), 1)

        retried = self.client.post(
            f"/api/v1/datasets/{self.dataset_id}/files",
            files={"file": ("recover.csv", b"value\n1\n", "text/csv")},
        )
        self.assertEqual(retried.status_code, 201, retried.text)
        self.assertEqual(retried.json()["file"]["status"], "Ready")
        self.assertEqual(len(self.storage.objects), 1)

    def test_database_finalization_failure_cleans_object_and_retry_recovers(self):
        class FailReadyCommitSession(Session):
            def commit(inner_self):
                if any(
                    isinstance(item, DatasetFile) and item.status == "Ready"
                    for item in inner_self.dirty
                ):
                    raise OperationalError("commit", {}, Exception())
                super().commit()

        failing_factory = sessionmaker(
            bind=self.engine,
            class_=FailReadyCommitSession,
            expire_on_commit=False,
        )
        with patch.object(main, "SessionLocal", failing_factory):
            response = self.client.post(
                f"/api/v1/datasets/{self.dataset_id}/files",
                files={"file": ("database.csv", b"value\n1\n", "text/csv")},
            )
        self.assertEqual(response.status_code, 500)
        self.assertEqual(self.storage.objects, {})
        session = self.session_factory()
        dataset_file = session.query(DatasetFile).one()
        self.assertEqual(dataset_file.status, "Processing")
        self.assertEqual(
            dataset_file.error_code,
            "metadata_persistence_failed",
        )
        self.assertEqual(session.get(Dataset, self.dataset_id).status, "Processing")
        session.close()

        retried = self.client.post(
            f"/api/v1/datasets/{self.dataset_id}/files",
            files={"file": ("database.csv", b"value\n1\n", "text/csv")},
        )
        self.assertEqual(retried.status_code, 201, retried.text)
        self.assertEqual(retried.json()["file"]["status"], "Ready")
        self.assertEqual(len(self.storage.objects), 1)

    def test_missing_object_on_replay_is_restored_from_verified_upload(self):
        content = b"value\n1\n"
        response = self.client.post(
            f"/api/v1/datasets/{self.dataset_id}/files",
            files={"file": ("missing.csv", content, "text/csv")},
        )
        self.assertEqual(response.status_code, 201, response.text)
        self.storage.objects.clear()

        replay = self.client.post(
            f"/api/v1/datasets/{self.dataset_id}/files",
            files={"file": ("missing.csv", content, "text/csv")},
        )
        self.assertEqual(replay.status_code, 201, replay.text)
        self.assertEqual(replay.json()["file"]["status"], "Ready")
        self.assertEqual(len(self.storage.objects), 1)
        session = self.session_factory()
        self.assertEqual(session.query(DatasetFile).count(), 1)
        self.assertEqual(session.query(DatasetFile).one().status, "Ready")
        self.assertEqual(session.get(Dataset, self.dataset_id).status, "Uploaded")
        session.close()

    def test_object_bytes_are_rechecked_against_checksum_on_replay(self):
        content = b"value\n1\n"
        response = self.client.post(
            f"/api/v1/datasets/{self.dataset_id}/files",
            files={"file": ("tampered.csv", content, "text/csv")},
        )
        self.assertEqual(response.status_code, 201, response.text)
        stored = next(iter(self.storage.objects.values()))
        stored["content"] = b"value\n2\n"

        replay = self.client.post(
            f"/api/v1/datasets/{self.dataset_id}/files",
            files={"file": ("tampered.csv", content, "text/csv")},
        )
        self.assertEqual(replay.status_code, 201, replay.text)
        self.assertEqual(len(self.storage.objects), 1)
        self.assertEqual(next(iter(self.storage.objects.values()))["content"], content)

    def test_legacy_success_record_replay_upgrades_integrity_metadata(self):
        content = b"value\n1\n"
        response = self.client.post(
            f"/api/v1/datasets/{self.dataset_id}/files",
            files={"file": ("legacy.csv", content, "text/csv")},
        )
        self.assertEqual(response.status_code, 201, response.text)
        file_id = response.json()["file"]["id"]

        session = self.session_factory()
        dataset_file = session.get(DatasetFile, file_id)
        dataset_file.idempotency_key = None
        session.commit()
        session.close()
        stored = next(iter(self.storage.objects.values()))
        stored["metadata"] = {}

        replay = self.client.post(
            f"/api/v1/datasets/{self.dataset_id}/files",
            files={"file": ("legacy.csv", content, "text/csv")},
        )
        self.assertEqual(replay.status_code, 201, replay.text)
        self.assertEqual(replay.json()["file"]["id"], file_id)
        self.assertEqual(len(self.storage.objects), 1)
        self.assertEqual(
            next(iter(self.storage.objects.values()))["metadata"]["sha256"],
            replay.json()["file"]["checksum"],
        )
        session = self.session_factory()
        self.assertEqual(session.query(DatasetFile).count(), 1)
        self.assertIsNotNone(session.get(DatasetFile, file_id).idempotency_key)
        session.close()

    def test_upload_for_missing_dataset_does_not_create_object_or_record(self):
        response = self.client.post(
            "/api/v1/datasets/999999/files",
            files={"file": ("missing.csv", b"value\n1\n", "text/csv")},
        )
        self.assertEqual(response.status_code, 404)
        self.assertEqual(self.storage.objects, {})
        session = self.session_factory()
        self.assertEqual(session.query(DatasetFile).count(), 0)
        session.close()

    def test_stage_24_type_validation_and_existing_health_endpoints(self):
        mismatch = self.client.post(
            f"/api/v1/datasets/{self.dataset_id}/files",
            files={"file": ("book.xlsx", b"content", "text/csv")},
        )
        self.assertEqual(mismatch.status_code, 415)

        self.assertEqual(self.client.get("/health").json()["status"], "healthy")
        with patch.object(main, "check_database", return_value=True):
            self.assertEqual(
                self.client.get("/health/database").json()["status"],
                "healthy",
            )
        with patch.object(main, "check_redis", return_value=True):
            self.assertEqual(
                self.client.get("/health/redis").json()["status"],
                "healthy",
            )
        with patch.object(main, "check_storage", return_value=(True, [])):
            self.assertEqual(
                self.client.get("/health/storage").json()["status"],
                "healthy",
            )
        with (
            patch.object(main, "check_database", return_value=True),
            patch.object(main, "check_redis", return_value=True),
            patch.object(main, "check_storage", return_value=(True, [])),
        ):
            self.assertEqual(
                self.client.get("/api/v1/system/info").json()["services"]["backend"],
                "healthy",
            )

    def test_dataset_creation_retrieval_listing_and_safe_deletion(self):
        created = self.client.post(
            f"/api/v1/projects/{self.project_id}/datasets",
            json={"name": "Created through API"},
        )
        self.assertEqual(created.status_code, 201, created.text)
        dataset_id = created.json()["id"]
        self.assertEqual(
            self.client.get(f"/api/v1/datasets/{dataset_id}").json()["status"],
            "Created",
        )
        listed = self.client.get(
            f"/api/v1/projects/{self.project_id}/datasets"
        )
        self.assertEqual(listed.status_code, 200)
        self.assertIn(dataset_id, [item["id"] for item in listed.json()])

        deleted = self.client.delete(f"/api/v1/datasets/{dataset_id}")
        self.assertEqual(deleted.status_code, 200)
        self.assertEqual(deleted.json()["deleted"], True)
        self.assertEqual(
            self.client.get(f"/api/v1/datasets/{dataset_id}").status_code,
            404,
        )

    def test_dataset_deletion_removes_its_raw_and_processed_objects_only(self):
        upload = self.client.post(
            f"/api/v1/datasets/{self.dataset_id}/files",
            files={"file": ("delete-me.csv", b"value\n1\n", "text/csv")},
        )
        self.assertEqual(upload.status_code, 201, upload.text)
        file_id = upload.json()["file"]["id"]
        normalized = self.client.post(
            f"/api/v1/datasets/{self.dataset_id}/files/{file_id}/normalize"
        )
        self.assertEqual(normalized.status_code, 200, normalized.text)

        other_dataset_id = self.dataset_id + 1000
        self.storage.objects[
            ("insightos-raw", f"{other_dataset_id}/other.csv")
        ] = {"content": b"untouched", "metadata": {}}
        self.storage.objects[
            ("insightos-processed", f"{other_dataset_id}/other.parquet")
        ] = {"content": b"untouched", "metadata": {}}

        response = self.client.delete(f"/api/v1/datasets/{self.dataset_id}")
        self.assertEqual(response.status_code, 200, response.text)
        self.assertFalse(
            any(
                key.startswith(f"{self.dataset_id}/")
                for bucket, key in self.storage.objects
                if bucket in {"insightos-raw", "insightos-processed"}
            )
        )
        self.assertIn(
            ("insightos-raw", f"{other_dataset_id}/other.csv"),
            self.storage.objects,
        )
        self.assertIn(
            ("insightos-processed", f"{other_dataset_id}/other.parquet"),
            self.storage.objects,
        )

    def test_dataset_deletion_storage_failure_keeps_database_records(self):
        upload = self.client.post(
            f"/api/v1/datasets/{self.dataset_id}/files",
            files={"file": ("keep-on-failure.csv", b"value\n1\n", "text/csv")},
        )
        self.assertEqual(upload.status_code, 201, upload.text)
        file_id = upload.json()["file"]["id"]
        normalized = self.client.post(
            f"/api/v1/datasets/{self.dataset_id}/files/{file_id}/normalize"
        )
        self.assertEqual(normalized.status_code, 200, normalized.text)
        objects_before = dict(self.storage.objects)
        original_remove = main._remove_dataset_objects

        def fail_processed(storage, bucket, dataset_id):
            if bucket == "insightos-processed":
                raise OSError("injected deletion failure")
            return original_remove(storage, bucket, dataset_id)

        with patch.object(main, "_remove_dataset_objects", side_effect=fail_processed):
            response = self.client.delete(
                f"/api/v1/datasets/{self.dataset_id}"
            )

        self.assertEqual(response.status_code, 502, response.text)
        session = self.session_factory()
        self.assertIsNotNone(session.get(Dataset, self.dataset_id))
        self.assertIsNotNone(session.get(DatasetFile, file_id))
        session.close()
        self.assertEqual(self.storage.objects, objects_before)

    def test_dataset_deletion_conflicts_with_active_normalization(self):
        upload = self.client.post(
            f"/api/v1/datasets/{self.dataset_id}/files",
            files={"file": ("busy.csv", b"value\n1\n", "text/csv")},
        )
        self.assertEqual(upload.status_code, 201, upload.text)
        file_id = upload.json()["file"]["id"]
        acquired = set()
        lock_ids = {}
        output_removal_started = Event()
        allow_deletion_to_continue = Event()
        original_remove = main._remove_dataset_objects

        def acquire(session, connection, _dataset_id, key):
            self.assertIs(session.get_bind(), connection)
            if key in acquired:
                return None
            acquired.add(key)
            lock_id = 999 + len(lock_ids)
            lock_ids[lock_id] = key
            return lock_id

        def release(_session, _connection, lock_id):
            acquired.discard(lock_ids.pop(lock_id))

        def pause_removal(storage, bucket, dataset_id):
            if bucket == "insightos-processed":
                output_removal_started.set()
                if not allow_deletion_to_continue.wait(timeout=10):
                    raise TimeoutError("test did not release dataset deletion")
            return original_remove(storage, bucket, dataset_id)

        with (
            patch.object(main, "_try_acquire_upload_lock", side_effect=acquire),
            patch.object(main, "_release_upload_lock", side_effect=release),
            patch.object(main, "_remove_dataset_objects", side_effect=pause_removal),
            ThreadPoolExecutor(max_workers=1) as executor,
        ):
            deletion_future = executor.submit(
                self.client.delete,
                f"/api/v1/datasets/{self.dataset_id}",
            )
            try:
                self.assertTrue(output_removal_started.wait(timeout=5))
                normalization = self.client.post(
                    f"/api/v1/datasets/{self.dataset_id}/files/{file_id}/normalize"
                )
                self.assertEqual(normalization.status_code, 409, normalization.text)
                self.assertEqual(
                    normalization.json()["detail"]["code"],
                    "normalization_in_progress",
                )
            finally:
                allow_deletion_to_continue.set()
            deleted = deletion_future.result(timeout=10)

        self.assertEqual(deleted.status_code, 200, deleted.text)
        self.assertFalse(
            any(
                key.startswith(f"{self.dataset_id}/")
                for bucket, key in self.storage.objects
                if bucket in {"insightos-raw", "insightos-processed"}
            )
        )


if __name__ == "__main__":
    unittest.main()
