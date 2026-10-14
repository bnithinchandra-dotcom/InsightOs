import base64
import datetime
import hashlib
import tempfile
import unittest
from decimal import Decimal
from io import BytesIO
from pathlib import Path
from unittest.mock import patch

import pyarrow as pa
import pyarrow.parquet as pq
from openpyxl import Workbook

from app.execution import ExecutionBudget
from app.normalization import (
    MAX_BATCH_SIZE,
    NormalizationError,
    _BudgetedOutput,
    _OutputLimitReached,
    _json_load,
    get_normalization_configuration,
    normalize_file,
    verify_normalization_result_detailed,
)
from app.parsing import ParsingError, parse_dataset_file


class FakeObjectStorage:
    def __init__(self):
        self.objects = {}
        self.put_count = 0
        self.fail_put = False

    def bucket_exists(self, bucket):
        return any(object_bucket == bucket for object_bucket, _key in self.objects)

    def make_bucket(self, _bucket):
        return None

    def stat_object(self, bucket, key):
        content = self.objects.get((bucket, key))
        return type("ObjectStat", (), {"size": len(content) if content is not None else -1})()

    def get_object(self, bucket, key):
        return BytesIO(self.objects[(bucket, key)])

    def remove_object(self, bucket, key):
        self.objects.pop((bucket, key), None)

    def put_object(
        self,
        bucket,
        key,
        stream,
        _size,
        content_type=None,
        metadata=None,
    ):
        if self.fail_put:
            raise OSError("injected storage failure")
        self.put_count += 1
        self.objects[(bucket, key)] = stream.read()


class NormalizationTests(unittest.TestCase):
    def test_budgeted_output_enforces_exact_cap_and_releases_failed_write_reservations(self):
        budget = ExecutionBudget({"normalized_output_bytes": 4})
        output = BytesIO()
        sink = _BudgetedOutput(output, budget)
        self.assertEqual(sink.write(b"1234"), 4)
        self.assertEqual(budget.used("normalized_output_bytes"), 4)
        with self.assertRaises(_OutputLimitReached):
            sink.write(b"5")
        self.assertEqual(output.getvalue(), b"1234")
        self.assertEqual(budget.used("normalized_output_bytes"), 4)

        class FailingBuffer(BytesIO):
            def write(self, _data):
                raise OSError("injected local write failure")

        failed_budget = ExecutionBudget({"normalized_output_bytes": 4})
        failed_sink = _BudgetedOutput(FailingBuffer(), failed_budget)
        with self.assertRaisesRegex(OSError, "injected local write failure"):
            failed_sink.write(b"1234")
        self.assertEqual(failed_budget.used("normalized_output_bytes"), 0)
        self.assertEqual(failed_budget.remaining("normalized_output_bytes"), 4)

    def normalize(self, filename: str, content: bytes, storage=None):
        storage = storage or FakeObjectStorage()
        parsed = parse_dataset_file(filename, BytesIO(content), len(content)).as_dict()
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / filename
            path.write_bytes(content)
            result = normalize_file(
                path,
                filename,
                parsed,
                dataset_id=7,
                file_id=11,
                source_checksum=hashlib.sha256(content).hexdigest(),
                storage=storage,
                processed_bucket="processed",
            )
        return result, storage

    def output_table(self, result, storage, table_index=0):
        table_result = result["tables"][table_index]
        output = storage.objects[("processed", table_result["object_key"])]
        return table_result, pq.read_table(BytesIO(output))

    def test_csv_tsv_keep_identifier_text_empty_strings_and_names(self):
        for filename, content in (
            ("people.csv", b"code,name\n00123,\n00007,Ada\n"),
            ("people.tsv", b"code\tname\n00123\t\n00007\tAda\n"),
        ):
            with self.subTest(filename=filename):
                result, storage = self.normalize(filename, content)
                table_result, table = self.output_table(result, storage)
                self.assertEqual(table.schema.names, ["code", "name"])
                self.assertEqual(table["code"].to_pylist(), ["00123", "00007"])
                self.assertEqual(table["name"].to_pylist(), ["", "Ada"])
                self.assertEqual(table_result["row_count"], 2)
                self.assertEqual(table_result["empty_value_counts"], [0, 1])

    def test_json_preserves_numeric_precision_null_missing_and_nested_values(self):
        content = (
            b'[{"amount":9007199254740993,"optional":null,"nested":{"x":[1,true]}},'
            b'{"amount":1.25,"nested":{"x":[1,true]}},'
            b'{"amount":2.50,"optional":"","nested":[1,2]}]'
        )
        result, storage = self.normalize("data.json", content)
        table_result, table = self.output_table(result, storage)
        self.assertEqual(table_result["row_count"], 3)
        self.assertEqual(str(table.schema.field("amount").type), "decimal256(18, 2)")
        self.assertEqual(
            table["amount"].to_pylist(),
            [
                Decimal("9007199254740993.00"),
                Decimal("1.25"),
                Decimal("2.50"),
            ],
        )
        self.assertEqual(
            table["optional"].to_pylist(),
            ['["null"]', '["missing"]', '["string",""]'],
        )
        self.assertEqual(table_result["null_value_counts"][1], 0)
        self.assertEqual(table_result["logical_null_value_counts"][1], 2)
        self.assertEqual(table_result["logical_empty_value_counts"][1], 1)
        encoded = table["nested"].to_pylist()
        self.assertTrue(all(value.startswith("[") for value in encoded))

    def test_json_preserves_missing_and_null_when_all_present_values_are_null(self):
        result, storage = self.normalize(
            "nulls.json",
            b'[{"value":null},{}]',
        )
        _, table = self.output_table(result, storage)
        self.assertEqual(
            table["value"].to_pylist(),
            ['["null"]', '["missing"]'],
        )

    def test_json_single_property_record_wrapper_is_supported(self):
        result, storage = self.normalize(
            "wrapped.json",
            b'{"records":[{"id":1},{"id":2}]}',
        )
        _, table = self.output_table(result, storage)
        self.assertEqual(table["id"].to_pylist(), [1, 2])

    def test_json_wrapper_siblings_and_duplicate_keys_are_rejected(self):
        for content in (
            b'{"records":[{"id":1}],"extra":[{"id":2}]}',
            b'{"records":[{"id":1}],"label":"kept nowhere"}',
            b'{"value":1,"value":2}',
        ):
            with self.subTest(content=content):
                with self.assertRaises(ParsingError):
                    self.normalize("ambiguous.json", content)

    def test_xlsx_keeps_worksheets_separate_and_preserves_dates(self):
        workbook = Workbook()
        first = workbook.active
        first.title = "Identifiers"
        first.append(["code", "active"])
        first.append(["001", True])
        second = workbook.create_sheet("Dates")
        second.append(["date"])
        second.append([datetime.date(2024, 1, 2)])
        stream = BytesIO()
        workbook.save(stream)

        result, storage = self.normalize("book.xlsx", stream.getvalue())
        self.assertEqual([table["table_name"] for table in result["tables"]], ["Identifiers", "Dates"])
        _, identifiers = self.output_table(result, storage, 0)
        _, dates = self.output_table(result, storage, 1)
        self.assertEqual(identifiers["code"].to_pylist(), ["001"])
        self.assertEqual(identifiers["active"].to_pylist(), [True])
        self.assertEqual(str(dates.schema.field("date").type), "date32[day]")
        self.assertEqual(dates["date"].to_pylist()[0].isoformat(), "2024-01-02")

    def test_xlsx_datetime_preserves_non_midnight_time(self):
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.append(["timestamp"])
        worksheet.append([datetime.datetime(2024, 2, 3, 16, 45, 12)])
        stream = BytesIO()
        workbook.save(stream)

        result, storage = self.normalize("timestamp.xlsx", stream.getvalue())
        _, table = self.output_table(result, storage)
        self.assertEqual(str(table.schema.field("timestamp").type), "timestamp[us]")
        self.assertEqual(
            table["timestamp"].to_pylist(),
            [datetime.datetime(2024, 2, 3, 16, 45, 12)],
        )

    def test_xlsx_mixed_date_and_datetime_uses_timestamp_without_truncation(self):
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.append(["when"])
        worksheet.append([datetime.date(2024, 2, 3)])
        worksheet.append([datetime.datetime(2024, 2, 4, 9, 8, 7)])
        stream = BytesIO()
        workbook.save(stream)

        result, storage = self.normalize("mixed-dates.xlsx", stream.getvalue())
        _, table = self.output_table(result, storage)
        self.assertEqual(str(table.schema.field("when").type), "timestamp[us]")
        self.assertEqual(
            table["when"].to_pylist(),
            [
                datetime.datetime(2024, 2, 3),
                datetime.datetime(2024, 2, 4, 9, 8, 7),
            ],
        )

    def test_xlsx_data_after_an_empty_first_row_is_not_silently_omitted(self):
        workbook = Workbook()
        blank_leading = workbook.active
        blank_leading.title = "Blank leading row"
        blank_leading["A2"] = "value"
        blank_leading["A3"] = 1
        valid = workbook.create_sheet("Valid")
        valid.append(["value"])
        valid.append([2])
        stream = BytesIO()
        workbook.save(stream)

        with self.assertRaises(NormalizationError) as error:
            self.normalize("leading-blank.xlsx", stream.getvalue())
        self.assertEqual(error.exception.code, "unsupported_structure")

    def test_empty_csv_and_empty_xlsx_worksheets_are_retained(self):
        result, storage = self.normalize("empty.csv", b"name,code\n")
        table_result, table = self.output_table(result, storage)
        self.assertEqual(table_result["row_count"], 0)
        self.assertEqual(table.schema.names, ["name", "code"])

        workbook = Workbook()
        workbook.active.title = "Blank"
        header_only = workbook.create_sheet("Header only")
        header_only.append(["name"])
        stream = BytesIO()
        workbook.save(stream)
        workbook_result, workbook_storage = self.normalize(
            "empty-sheets.xlsx",
            stream.getvalue(),
        )
        self.assertEqual(
            [table["table_name"] for table in workbook_result["tables"]],
            ["Blank", "Header only"],
        )
        self.assertIsNone(workbook_result["tables"][0]["object_key"])
        header_table, output = self.output_table(
            workbook_result,
            workbook_storage,
            1,
        )
        self.assertEqual(header_table["row_count"], 0)
        self.assertEqual(output.schema.names, ["name"])

    def test_xml_keeps_flattened_namespace_paths_missing_and_empty(self):
        content = (
            b'<d:root xmlns:d="urn:records"><d:record id="1">'
            b"<d:value></d:value></d:record><d:record id=\"2\"/>"
            b"</d:root>"
        )
        result, storage = self.normalize("data.xml", content)
        _, table = self.output_table(result, storage)
        self.assertEqual(table.schema.names, ["@id", "{urn:records}value"])
        self.assertEqual(table["@id"].to_pylist(), ["1", "2"])
        self.assertEqual(table["{urn:records}value"].to_pylist(), ["", None])

    def test_xml_preserves_root_container_and_record_attributes(self):
        content = (
            b'<d:root xmlns:d="urn:records" batch="B1">'
            b'<d:records version="v2">'
            b'<d:record id="1"><d:value>A</d:value></d:record>'
            b'<d:record id="2"><d:value>B</d:value></d:record>'
            b"</d:records></d:root>"
        )
        result, storage = self.normalize("attributes.xml", content)
        _, table = self.output_table(result, storage)

        self.assertEqual(
            table.schema.names,
            [
                "@id",
                "{urn:records}value",
                '@container:["{urn:records}root"]/@batch',
                '@container:["{urn:records}root","{urn:records}records"]/@version',
            ],
        )
        self.assertEqual(table["@id"].to_pylist(), ["1", "2"])
        self.assertEqual(
            table['@container:["{urn:records}root"]/@batch'].to_pylist(),
            ["B1", "B1"],
        )
        self.assertEqual(
            table[
                '@container:["{urn:records}root","{urn:records}records"]/@version'
            ].to_pylist(),
            ["v2", "v2"],
        )

    def test_existing_parquet_schema_and_values_are_preserved(self):
        source = BytesIO()
        pq.write_table(
            pa.table(
                {
                    "code": pa.array(["001", None], type=pa.string()),
                    "value": pa.array([1.25, 2.5], type=pa.float64()),
                }
            ),
            source,
        )
        result, storage = self.normalize("source.parquet", source.getvalue())
        _, output = self.output_table(result, storage)
        self.assertEqual(output.schema.names, ["code", "value"])
        self.assertEqual(output["code"].to_pylist(), ["001", None])
        self.assertEqual(output["value"].to_pylist(), [1.25, 2.5])

    def test_same_source_and_configuration_reuse_stable_output_key(self):
        content = b"id\n1\n2\n"
        storage = FakeObjectStorage()
        first, _ = self.normalize("retry.csv", content, storage)
        first_key = first["tables"][0]["object_key"]
        first_puts = storage.put_count
        second, _ = self.normalize("retry.csv", content, storage)
        self.assertEqual(second["tables"][0]["object_key"], first_key)
        self.assertEqual(storage.put_count, first_puts)

    def test_storage_failure_is_explicit_and_does_not_mutate_source(self):
        content = b"value\n1\n"
        storage = FakeObjectStorage()
        storage.fail_put = True
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "failure.csv"
            path.write_bytes(content)
            parsed = parse_dataset_file("failure.csv", BytesIO(content), len(content)).as_dict()
            with self.assertRaises(OSError):
                normalize_file(
                    path,
                    "failure.csv",
                    parsed,
                    1,
                    2,
                    hashlib.sha256(content).hexdigest(),
                    storage,
                    "processed",
                )
            self.assertEqual(path.read_bytes(), content)

    def test_direct_normalization_rejects_source_checksum_mismatch_before_writing(self):
        content = b"value\n1\n"
        storage = FakeObjectStorage()
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "source.csv"
            path.write_bytes(content)
            parsed = parse_dataset_file(
                "source.csv",
                BytesIO(content),
                len(content),
            ).as_dict()
            with self.assertRaises(NormalizationError) as error:
                normalize_file(
                    path,
                    "source.csv",
                    parsed,
                    7,
                    11,
                    "0" * 64,
                    storage,
                    "processed",
                )
        self.assertEqual(error.exception.code, "source_integrity_failed")
        self.assertEqual(storage.put_count, 0)

    def test_detailed_normalization_verification_reports_integrity_and_provenance(self):
        result, storage = self.normalize("verified.csv", b"value\n1\n")
        verification = verify_normalization_result_detailed(
            storage,
            "processed",
            result,
        )
        self.assertEqual(verification["status"], "verified")
        self.assertTrue(
            all(check["status"] == "verified" for check in verification["checks"])
        )
        self.assertTrue(verification["limitations"])

        corrupted = dict(result)
        corrupted["provenance"] = dict(result["provenance"])
        corrupted["provenance"]["outputs"] = []
        stale_verification = verify_normalization_result_detailed(
            storage,
            "processed",
            corrupted,
        )
        self.assertEqual(stale_verification["status"], "failed")
        self.assertEqual(
            stale_verification["checks"],
            [{"name": "normalization_metadata", "status": "failed"}],
        )

        malformed = dict(result)
        malformed["tables"] = ["not-table-metadata"]
        malformed_verification = verify_normalization_result_detailed(
            storage,
            "processed",
            malformed,
        )
        self.assertEqual(malformed_verification["status"], "failed")

    def test_output_size_limit_is_enforced(self):
        values = [
            base64.b64encode(
                hashlib.shake_256(f"stable-normalization-limit-fixture-{index}".encode())
                .digest(60_000)
            )
            for index in range(20)
        ]
        content = b"value\n" + b"\n".join(values) + b"\n"
        storage = FakeObjectStorage()
        with patch.dict("os.environ", {"MAX_NORMALIZED_OUTPUT_MB": "1"}):
            with self.assertRaises(NormalizationError) as error:
                self.normalize("large.csv", content, storage)
        self.assertEqual(error.exception.code, "resource_limit")
        self.assertEqual(storage.put_count, 0)

    def test_output_limit_during_writer_cleanup_preserves_typed_failure(self):
        created_paths = []
        original = tempfile.NamedTemporaryFile

        def tracking_temp_file(*args, **kwargs):
            output = original(*args, **kwargs)
            created_paths.append(Path(output.name))
            return output

        with (
            patch(
                "app.normalization._normalization_settings",
                return_value=("zstd", 8192, 10, "0" * 64),
            ),
            patch(
                "app.normalization.tempfile.NamedTemporaryFile",
                side_effect=tracking_temp_file,
            ),
            self.assertRaises(NormalizationError) as error,
        ):
            self.normalize("tiny-output-cap.csv", b"value\nx\n")

        self.assertEqual(error.exception.code, "resource_limit")
        self.assertEqual(len(created_paths), 1)
        self.assertFalse(created_paths[0].exists())

    def test_output_size_limit_removes_earlier_table_objects_from_failed_operation(self):
        workbook = Workbook()
        first = workbook.active
        first.title = "small"
        first.append(["value"])
        first.append(["kept only if all tables normalize"])
        second = workbook.create_sheet("large")
        second.append(["value"])
        for index in range(60):
            value = base64.b64encode(
                hashlib.shake_256(
                    f"stable-multitable-limit-fixture-{index}".encode()
                ).digest(60_000)
            ).decode("ascii")
            second.append([value])
        output = BytesIO()
        workbook.save(output)
        storage = FakeObjectStorage()

        with (
            patch.dict("os.environ", {"MAX_NORMALIZED_OUTPUT_MB": "1"}),
            self.assertRaises(NormalizationError) as error,
        ):
            self.normalize("multitable.xlsx", output.getvalue(), storage)

        self.assertEqual(error.exception.code, "resource_limit")
        self.assertEqual(storage.put_count, 1)
        self.assertEqual(storage.objects, {})

    def test_failed_output_validation_removes_the_temporary_parquet_file(self):
        content = b"value\none\ntwo\n"
        created_paths = []
        original = tempfile.NamedTemporaryFile

        def tracking_temp_file(*args, **kwargs):
            output = original(*args, **kwargs)
            created_paths.append(Path(output.name))
            return output

        with (
            patch(
                "app.normalization.tempfile.NamedTemporaryFile",
                side_effect=tracking_temp_file,
            ),
            patch(
                "app.normalization._validate_parquet",
                side_effect=NormalizationError(
                    "fidelity_validation_failed",
                    "Injected validation failure.",
                ),
            ),
            self.assertRaises(NormalizationError),
        ):
            self.normalize("validation-failure.csv", content)

        self.assertEqual(len(created_paths), 1)
        self.assertFalse(created_paths[0].exists())

    def test_json_normalization_repeats_source_and_depth_preflight(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "input.json"
            path.write_bytes(b'{"records":[]}' + b" " * (16 * 1024))
            with (
                patch.dict("os.environ", {"MAX_JSON_ESTIMATED_DECODED_MB": "1"}),
                patch("app.normalization.json.loads") as loads,
                self.assertRaises(NormalizationError) as error,
            ):
                _json_load(path)
            self.assertEqual(error.exception.code, "resource_limit")
            loads.assert_not_called()

            path.write_text('{"a":{"b":{"c":1}}}', encoding="utf-8")
            with (
                patch.dict("os.environ", {"MAX_JSON_DEPTH": "2"}),
                patch("app.normalization.json.loads") as loads,
                self.assertRaises(NormalizationError) as error,
            ):
                _json_load(path)
            self.assertEqual(error.exception.code, "resource_limit")
            loads.assert_not_called()

    def test_normalization_batch_size_has_a_hard_upper_bound(self):
        with patch.dict("os.environ", {"NORMALIZATION_BATCH_SIZE": "0"}):
            with self.assertRaises(NormalizationError) as error:
                get_normalization_configuration()
        self.assertEqual(error.exception.code, "invalid_configuration")

        with patch.dict(
            "os.environ",
            {"NORMALIZATION_BATCH_SIZE": str(MAX_BATCH_SIZE + 1)},
        ):
            with self.assertRaises(NormalizationError) as error:
                get_normalization_configuration()
        self.assertEqual(error.exception.code, "invalid_configuration")


if __name__ == "__main__":
    unittest.main()
