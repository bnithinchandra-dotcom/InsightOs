import csv
import hashlib
import json
import math
import os
import tempfile
from dataclasses import dataclass
from datetime import date, datetime, time
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Iterable, Iterator, Sequence

from minio.error import MinioException, S3Error
import pyarrow as pa
import pyarrow.parquet as pq
from openpyxl import load_workbook
from urllib3.exceptions import HTTPError

from app.parsing import (
    MISSING,
    ParserConfigurationError,
    ParsingError,
    _flatten_xml_record,
    _parse_xml_tree,
    _xml_candidates,
    _xml_depths,
    get_parse_limits,
    parse_dataset_file,
)


NORMALIZATION_PIPELINE_VERSION = "1"
DEFAULT_BATCH_SIZE = 8192
MAX_BATCH_SIZE = 65536
DEFAULT_MAX_OUTPUT_BYTES = 600 * 1024 * 1024
_MAX_JSON_DECIMAL_PRECISION = 76


class NormalizationError(Exception):
    def __init__(self, code: str, message: str):
        super().__init__(message)
        self.code = code
        self.message = message


@dataclass
class LogicalTable:
    name: str | None
    columns: list[dict]
    rows: Iterable[Sequence]
    column_values: Sequence[Sequence] | None = None


@dataclass
class _JsonColumnValues(Sequence):
    records: list[dict]
    name: str

    def __len__(self) -> int:
        return len(self.records)

    def __getitem__(self, index):
        return self.records[index].get(self.name)


def _positive_env(name: str, default: int) -> int:
    supplied = os.getenv(name)
    try:
        value = int(supplied) if supplied is not None else default
    except ValueError:
        raise NormalizationError(
            "invalid_configuration",
            f"{name} must be a positive integer.",
        ) from None
    if value <= 0:
        raise NormalizationError(
            "invalid_configuration",
            f"{name} must be a positive integer.",
        )
    return value


def _normalization_settings() -> tuple[str, int, int, str]:
    compression_value = os.getenv("NORMALIZATION_COMPRESSION", "zstd").strip().lower()
    compression = None if compression_value in {"none", "uncompressed"} else compression_value
    if compression is not None:
        try:
            available = pa.Codec.is_available(compression)
        except ValueError:
            available = False
        if not available:
            raise NormalizationError(
                "invalid_configuration",
                "NORMALIZATION_COMPRESSION must name an available Parquet codec.",
            )
    batch_size = _positive_env("NORMALIZATION_BATCH_SIZE", DEFAULT_BATCH_SIZE)
    if batch_size > MAX_BATCH_SIZE:
        raise NormalizationError(
            "invalid_configuration",
            f"NORMALIZATION_BATCH_SIZE must not exceed {MAX_BATCH_SIZE}.",
        )
    max_output_bytes = (
        _positive_env("MAX_NORMALIZED_OUTPUT_MB", 600) * 1024 * 1024
    )
    config = {
        "pipeline_version": NORMALIZATION_PIPELINE_VERSION,
        "compression": compression or "uncompressed",
        "batch_size": batch_size,
        "max_output_bytes": max_output_bytes,
    }
    config_hash = hashlib.sha256(
        json.dumps(config, sort_keys=True, separators=(",", ":")).encode("utf-8")
    ).hexdigest()
    return compression or "uncompressed", batch_size, max_output_bytes, config_hash


def _json_load(path: Path):
    def reject_constant(_value):
        raise NormalizationError(
            "invalid_source",
            "The JSON source contains a non-standard numeric value.",
        )

    def reject_duplicate_keys(pairs):
        result = {}
        for key, value in pairs:
            if key in result:
                raise NormalizationError(
                    "invalid_source",
                    "The JSON source contains a duplicate object key.",
                )
            result[key] = value
        return result

    try:
        with path.open("r", encoding="utf-8-sig", newline="") as source:
            return json.load(
                source,
                object_pairs_hook=reject_duplicate_keys,
                parse_float=Decimal,
                parse_int=int,
                parse_constant=reject_constant,
            )
    except NormalizationError:
        raise
    except (OSError, UnicodeError, json.JSONDecodeError, RecursionError):
        raise NormalizationError(
            "invalid_source",
            "The JSON source could not be read using its supported structure.",
        ) from None


def _json_records(document):
    if isinstance(document, list):
        return document
    if isinstance(document, dict):
        arrays = [
            value
            for value in document.values()
            if isinstance(value, list)
            and all(isinstance(item, dict) for item in value)
        ]
        if arrays:
            if len(document) != 1 or len(arrays) != 1:
                raise NormalizationError(
                    "unsupported_structure",
                    "JSON wrapper siblings are not supported for normalization.",
                )
            return arrays[0]
        if all(not isinstance(value, (dict, list)) for value in document.values()):
            return [document]
    raise NormalizationError(
        "unsupported_structure",
        "JSON normalization supports an array of objects, a flat object, or a "
        "single-property object-array wrapper.",
    )


def _json_encode(value):
    if value is None:
        return ["null"]
    if isinstance(value, bool):
        return ["boolean", value]
    if isinstance(value, int):
        return ["integer", str(value)]
    if isinstance(value, Decimal):
        return ["number", str(value)]
    if isinstance(value, float):
        if not math.isfinite(value):
            raise NormalizationError(
                "unsupported_value",
                "Non-finite numeric values cannot be normalized safely.",
            )
        return ["number", value.hex()]
    if isinstance(value, str):
        return ["string", value]
    if isinstance(value, list):
        return ["array", [_json_encode(item) for item in value]]
    if isinstance(value, dict):
        return [
            "object",
            [
                [str(key), _json_encode(child)]
                for key, child in sorted(value.items(), key=lambda item: str(item[0]))
            ],
        ]
    raise NormalizationError(
        "unsupported_value",
        "A JSON value could not be represented without loss.",
    )


def _typed_cell_encode(value):
    if value is None:
        return None
    if isinstance(value, datetime):
        tag = "datetime"
        encoded = value.isoformat()
    elif isinstance(value, date):
        tag = "date"
        encoded = value.isoformat()
    elif isinstance(value, time):
        tag = "time"
        encoded = value.isoformat()
    elif isinstance(value, bool):
        tag = "boolean"
        encoded = value
    elif isinstance(value, int):
        tag = "integer"
        encoded = str(value)
    elif isinstance(value, float):
        if not math.isfinite(value):
            raise NormalizationError(
                "unsupported_value",
                "Non-finite numeric values cannot be normalized safely.",
            )
        tag = "number"
        encoded = value.hex()
    elif isinstance(value, Decimal):
        tag = "decimal"
        encoded = str(value)
    elif isinstance(value, str):
        tag = "string"
        encoded = value
    else:
        raise NormalizationError(
            "unsupported_value",
            "A cell value could not be represented without loss.",
        )
    return json.dumps([tag, encoded], ensure_ascii=False, separators=(",", ":"))


def _read_csv_table(path: Path, detected_format: str, parsed_columns: list[dict]):
    delimiter = "\t" if detected_format == "tsv" else ","

    def rows() -> Iterator[list[str]]:
        try:
            with path.open(
                "r",
                encoding="utf-8-sig",
                errors="strict",
                newline="",
            ) as source:
                reader = csv.reader(source, delimiter=delimiter, strict=True)
                next(reader, None)
                for row in reader:
                    yield row
        except (OSError, UnicodeError, csv.Error):
            raise NormalizationError(
                "invalid_source",
                "Delimited source could not be reread consistently for normalization.",
            ) from None

    yield LogicalTable(None, parsed_columns, rows())


def _read_xlsx_tables(path: Path, parsed: dict):
    workbook = load_workbook(
        path,
        read_only=True,
        data_only=False,
        keep_links=False,
    )
    try:
        worksheet_metadata = parsed["metadata"].get("worksheets", [])
        if len(worksheet_metadata) != len(workbook.worksheets):
            raise NormalizationError(
                "invalid_source",
                "The XLSX worksheet list changed after parsing.",
            )
        for worksheet, metadata in zip(workbook.worksheets, worksheet_metadata):
            columns = metadata["columns"]
            iterator = worksheet.iter_rows()
            if not columns:
                try:
                    has_values = any(
                        cell.value is not None
                        for row in iterator
                        for cell in row
                    )
                finally:
                    close_iterator = getattr(iterator, "close", None)
                    if close_iterator is not None:
                        close_iterator()
                if has_values:
                    raise NormalizationError(
                        "unsupported_structure",
                        "An XLSX worksheet contains values after an empty first row "
                        "and cannot be normalized without dropping them.",
                    )
                yield LogicalTable(metadata["name"], columns, ())
                continue
            next(iterator, None)

            def rows(row_iterator=iterator, expected=len(columns)):
                for cells in row_iterator:
                    if len(cells) != expected:
                        raise NormalizationError(
                            "invalid_source",
                            "An XLSX worksheet row no longer matches its parsed schema.",
                        )
                    values = []
                    for cell in cells:
                        if cell.data_type == "e":
                            raise NormalizationError(
                                "unsupported_value",
                                "XLSX error cells cannot be normalized without changing their meaning.",
                            )
                        values.append(cell.value)
                    yield values

            yield LogicalTable(metadata["name"], columns, rows())
    except (OSError, ValueError, KeyError) as error:
        if isinstance(error, NormalizationError):
            raise
        raise NormalizationError(
            "invalid_source",
            "The XLSX source could not be reread consistently for normalization.",
        ) from None
    finally:
        workbook.close()


def _read_json_tables(path: Path, parsed: dict):
    document = _json_load(path)
    records = _json_records(document)
    columns = parsed["columns"]

    def rows():
        for record in records:
            if not isinstance(record, dict):
                raise NormalizationError(
                    "unsupported_structure",
                    "Every normalized JSON record must be an object.",
                )
            yield [record.get(column["name"], MISSING) for column in columns]

    yield LogicalTable(
        None,
        columns,
        rows(),
        [_JsonColumnValues(records, column["name"]) for column in columns],
    )


def _read_parquet_tables(path: Path, parsed: dict):
    parquet_file = pq.ParquetFile(path)
    columns = parsed["columns"]

    def rows():
        try:
            for batch in parquet_file.iter_batches(batch_size=8192):
                for index in range(batch.num_rows):
                    yield [
                        batch.column(column_index)[index].as_py()
                        for column_index in range(batch.num_columns)
                    ]
        finally:
            parquet_file.close()

    yield LogicalTable(None, columns, rows())


def _read_xml_tables(path: Path, parsed: dict):
    with path.open("rb") as source:
        root, _element_count = _parse_xml_tree(source, get_parse_limits())
    depths, _count = _xml_depths(root, get_parse_limits())
    candidates = _xml_candidates(root, depths)
    if len(candidates) != 1:
        raise NormalizationError(
            "unsupported_structure",
            "The XML source no longer has the single record group selected during parsing.",
        )
    _depth, _parent, _tag, records = candidates[0]
    container_columns, container_values = _xml_container_fields(parsed)
    record_columns = parsed["columns"]
    record_names = [column["name"] for column in record_columns]
    container_names = [column["name"] for column in container_columns]
    if set(record_names).intersection(container_names):
        raise NormalizationError(
            "unsupported_structure",
            "XML container attributes conflict with record column names and "
            "cannot be represented without ambiguity.",
        )
    columns = record_columns + container_columns

    def rows():
        for record in records:
            try:
                flattened = _flatten_xml_record(record)
            except ParsingError as error:
                raise NormalizationError(error.code, error.message) from None
            yield [flattened.get(name) for name in record_names] + container_values

    yield LogicalTable(None, columns, rows())


def _xml_container_fields(parsed: dict) -> tuple[list[dict], list[str]]:
    ancestors = parsed.get("metadata", {}).get("container_attributes", [])
    if not isinstance(ancestors, list):
        raise NormalizationError(
            "invalid_source",
            "XML container attribute metadata is invalid.",
        )

    path = []
    columns = []
    values = []
    for ancestor in ancestors:
        if (
            not isinstance(ancestor, dict)
            or not isinstance(ancestor.get("element"), str)
            or not isinstance(ancestor.get("attributes"), dict)
            or any(
                not isinstance(name, str) or not isinstance(value, str)
                for name, value in ancestor["attributes"].items()
            )
        ):
            raise NormalizationError(
                "invalid_source",
                "XML container attribute metadata is invalid.",
            )
        path.append(ancestor["element"])
        path_name = json.dumps(path, ensure_ascii=False, separators=(",", ":"))
        for attribute, value in ancestor["attributes"].items():
            name = f"@container:{path_name}/@{attribute}"
            columns.append(
                {
                    "name": name,
                    "position": len(columns),
                    "physical_type": "string",
                    "physical_types": ["string"],
                    "missing_values": 0,
                    "empty_values": 0,
                }
            )
            values.append(value)
    return columns, values


def _table_sources(path: Path, parsed: dict):
    detected_format = parsed["detected_format"]
    if detected_format in {"csv", "tsv"}:
        yield from _read_csv_table(path, detected_format, parsed["columns"])
    elif detected_format == "xlsx":
        yield from _read_xlsx_tables(path, parsed)
    elif detected_format == "json":
        yield from _read_json_tables(path, parsed)
    elif detected_format == "parquet":
        yield from _read_parquet_tables(path, parsed)
    elif detected_format == "xml":
        yield from _read_xml_tables(path, parsed)
    else:
        raise NormalizationError(
            "unsupported_format",
            "This source format is not supported for normalization.",
        )


def _decimal_schema(values: Iterable) -> pa.DataType | None:
    max_scale = 0
    max_integer_digits = 0
    for value in values:
        if value is None:
            continue
        try:
            number = value if isinstance(value, Decimal) else Decimal(value)
        except (InvalidOperation, ValueError, TypeError):
            return None
        sign, digits, exponent = number.as_tuple()
        scale = max(0, -exponent)
        integer_digits = max(0, len(digits) + exponent)
        max_scale = max(max_scale, scale)
        max_integer_digits = max(max_integer_digits, integer_digits)
    precision = max_integer_digits + max_scale
    if max_scale > _MAX_JSON_DECIMAL_PRECISION or precision > _MAX_JSON_DECIMAL_PRECISION:
        return None
    precision = max(precision, max_scale, 1)
    return pa.decimal256(precision, max_scale)


def _arrow_type(
    detected_format: str,
    column: dict,
    values: Sequence | None,
    warnings: list[str],
) -> tuple[pa.DataType, str]:
    if detected_format in {"csv", "tsv", "xml"}:
        return pa.string(), "identity"
    if detected_format == "parquet":
        raise AssertionError("Parquet schema is handled separately.")

    physical_types = {
        value
        for value in column.get("physical_types", [])
        if value != "null"
    }
    if detected_format == "json":
        if (
            "null" in column.get("physical_types", [])
            and column.get("missing_values", 0) > 0
        ):
            warnings.append(
                f"Missing fields and explicit nulls in JSON column {column['name']!r} use reversible tagged JSON text."
            )
            return pa.string(), "tagged_json"
        if not physical_types:
            return pa.null(), "identity"
        if physical_types == {"integer"}:
            if values and all(
                value is None
                or isinstance(value, int)
                and not isinstance(value, bool)
                and -(2**63) <= value < 2**63
                for value in values
            ):
                return pa.int64(), "identity"
            data_type = _decimal_schema(values or ())
            if data_type is not None:
                return data_type, "decimal"
            warnings.append(
                f"Column {column['name']!r} uses reversible tagged JSON text because its numeric precision exceeds Parquet Decimal limits."
            )
            return pa.string(), "tagged_json"
        if physical_types <= {"integer", "number"}:
            data_type = _decimal_schema(values or ())
            if data_type is not None:
                return data_type, "decimal"
            warnings.append(
                f"Column {column['name']!r} uses reversible tagged JSON text because its numeric precision exceeds Parquet Decimal limits."
            )
            return pa.string(), "tagged_json"
        if physical_types == {"boolean"}:
            return pa.bool_(), "identity"
        if physical_types == {"string"}:
            return pa.string(), "identity"
        if physical_types <= {"object", "array"}:
            warnings.append(
                f"Nested JSON values in column {column['name']!r} use reversible tagged JSON text."
            )
            return pa.string(), "tagged_json"
        warnings.append(
            f"Mixed JSON values in column {column['name']!r} use reversible tagged JSON text."
        )
        return pa.string(), "tagged_json"

    if not physical_types:
        return pa.null(), "identity"

    if physical_types == {"integer"}:
        return pa.int64(), "identity"
    if physical_types == {"number"}:
        return pa.float64(), "identity"
    if physical_types == {"boolean"}:
        return pa.bool_(), "identity"
    if physical_types == {"string"}:
        return pa.string(), "identity"
    if physical_types == {"date"}:
        return pa.date32(), "identity"
    if physical_types == {"datetime"}:
        return pa.timestamp("us"), "identity"
    if physical_types <= {"date", "datetime"}:
        return pa.timestamp("us"), "datetime"
    if physical_types == {"time"}:
        return pa.time64("us"), "identity"
    if physical_types == {"formula"}:
        warnings.append(
            f"Formula expressions in column {column['name']!r} are preserved as text and are not evaluated."
        )
        return pa.string(), "identity"
    if "formula" in physical_types:
        raise NormalizationError(
            "unsupported_value",
            f"Column {column['name']!r} mixes formula cells with other values; "
            "normalization cannot preserve their distinction safely.",
        )

    warnings.append(
        f"Mixed XLSX values in column {column['name']!r} use reversible typed JSON text."
    )
    return pa.string(), "typed_json"


def _convert(value, encoding: str):
    if value is MISSING:
        return (
            json.dumps(["missing"], separators=(",", ":"))
            if encoding == "tagged_json"
            else None
        )
    if value is None:
        return (
            json.dumps(["null"], separators=(",", ":"))
            if encoding == "tagged_json"
            else None
        )
    if encoding == "decimal":
        return value if isinstance(value, Decimal) else Decimal(value)
    if encoding == "datetime" and isinstance(value, date) and not isinstance(value, datetime):
        return datetime.combine(value, datetime.min.time())
    if encoding == "tagged_json":
        return json.dumps(
            _json_encode(value),
            ensure_ascii=False,
            separators=(",", ":"),
        )
    if encoding == "typed_json":
        return _typed_cell_encode(value)
    return value


def _row_digest_update(digest, values: Sequence) -> None:
    def canonical(value):
        if isinstance(value, float):
            return ["float", value.hex()]
        if isinstance(value, Decimal):
            return ["decimal", str(value)]
        if isinstance(value, datetime):
            return ["datetime", value.isoformat()]
        if isinstance(value, date):
            return ["date", value.isoformat()]
        if isinstance(value, time):
            return ["time", value.isoformat()]
        if isinstance(value, bytes):
            return ["bytes", value.hex()]
        if isinstance(value, list):
            return ["list", [canonical(item) for item in value]]
        if isinstance(value, tuple):
            return ["tuple", [canonical(item) for item in value]]
        if isinstance(value, dict):
            return [
                "dict",
                [
                    [str(key), canonical(child)]
                    for key, child in sorted(value.items(), key=lambda item: str(item[0]))
                ],
            ]
        return value

    encoded = json.dumps(
        [canonical(value) for value in values],
        ensure_ascii=False,
        sort_keys=True,
        separators=(",", ":"),
        allow_nan=False,
    )
    digest.update(encoded.encode("utf-8"))
    digest.update(b"\n")


def _validate_parquet(
    path: Path,
    expected_schema: pa.Schema,
    expected_rows: int,
    expected_digest: str,
    expected_nulls: list[int],
    expected_empty: list[int],
) -> dict:
    parquet_file = None
    try:
        parquet_file = pq.ParquetFile(path)
        actual_schema = parquet_file.schema_arrow
        if actual_schema.names != expected_schema.names:
            raise NormalizationError(
                "fidelity_validation_failed",
                "Parquet read-back changed the normalized column names or order.",
            )
        if len(actual_schema) != len(expected_schema) or any(
            actual.type != expected.type
            for actual, expected in zip(actual_schema, expected_schema)
        ):
            raise NormalizationError(
                "fidelity_validation_failed",
                "Parquet read-back changed the normalized Arrow types.",
            )
        rows = 0
        null_counts = [0] * len(expected_schema)
        empty_counts = [0] * len(expected_schema)
        digest = hashlib.sha256()
        for batch in parquet_file.iter_batches(batch_size=DEFAULT_BATCH_SIZE):
            rows += batch.num_rows
            for row_index in range(batch.num_rows):
                values = [
                    batch.column(column_index)[row_index].as_py()
                    for column_index in range(batch.num_columns)
                ]
                _row_digest_update(digest, values)
                for column_index, value in enumerate(values):
                    if value is None:
                        null_counts[column_index] += 1
                    elif value == "":
                        empty_counts[column_index] += 1
        if (
            rows != expected_rows
            or digest.hexdigest() != expected_digest
            or null_counts != expected_nulls
            or empty_counts != expected_empty
        ):
            raise NormalizationError(
                "fidelity_validation_failed",
                "Parquet read-back did not match the normalized source values.",
            )
        return {
            "row_count": rows,
            "column_count": len(actual_schema),
            "null_value_counts": null_counts,
            "empty_value_counts": empty_counts,
            "schema": [
                {"name": field.name, "type": str(field.type)}
                for field in actual_schema
            ],
            "rows_sha256": digest.hexdigest(),
        }
    except NormalizationError:
        raise
    except (OSError, ValueError, pa.ArrowException):
        raise NormalizationError(
            "fidelity_validation_failed",
            "The generated Parquet file could not be read back reliably.",
        ) from None
    finally:
        if parquet_file is not None:
            parquet_file.close()


def _normalize_table(
    table: LogicalTable,
    source_path: Path,
    detected_format: str,
    dataset_id: int,
    file_id: int,
    source_checksum: str,
    config_hash: str,
    table_index: int,
    compression: str,
    batch_size: int,
    remaining_output_bytes: int,
    storage,
    processed_bucket: str,
) -> tuple[dict, int]:
    warnings: list[str] = []
    column_values = table.column_values

    if detected_format == "parquet":
        with pq.ParquetFile(source_path) as parquet_file:
            input_schema = parquet_file.schema_arrow
        arrow_schema = input_schema
        encodings = ["identity"] * len(arrow_schema)
    else:
        arrow_fields = []
        encodings = []
        for index, column in enumerate(table.columns):
            values = column_values[index] if column_values is not None else None
            arrow_type, encoding = _arrow_type(
                detected_format,
                column,
                values,
                warnings,
            )
            arrow_fields.append(pa.field(column["name"], arrow_type))
            encodings.append(encoding)
        arrow_schema = pa.schema(arrow_fields)

    metadata = {
        b"insightos.pipeline_version": NORMALIZATION_PIPELINE_VERSION.encode(),
        b"insightos.source_sha256": source_checksum.encode(),
        b"insightos.configuration_sha256": config_hash.encode(),
        b"insightos.detected_format": detected_format.encode(),
        b"insightos.table_index": str(table_index).encode(),
        b"insightos.table_name": (table.name or "").encode("utf-8"),
    }
    arrow_schema = arrow_schema.with_metadata(
        {**(arrow_schema.metadata or {}), **metadata}
    )
    output_temp = tempfile.NamedTemporaryFile(prefix="insightos-normalized-", suffix=".parquet", delete=False)
    output_path = Path(output_temp.name)
    output_temp.close()
    expected_digest = hashlib.sha256()
    expected_nulls = [0] * len(arrow_schema)
    expected_empty = [0] * len(arrow_schema)
    row_count = 0
    writer = None

    try:
        def check_temporary_size() -> None:
            if output_path.stat().st_size > remaining_output_bytes:
                raise NormalizationError(
                    "resource_limit",
                    "Normalized output exceeds MAX_NORMALIZED_OUTPUT_MB.",
                )

        try:
            writer = pq.ParquetWriter(
                output_path,
                arrow_schema,
                compression=None if compression == "uncompressed" else compression,
            )
            current_rows = []
            try:
                for row in table.rows:
                    if len(row) != len(arrow_schema):
                        raise NormalizationError(
                            "invalid_source",
                            "A source row no longer matches its parsed column count.",
                        )
                    current_rows.append(
                        [
                            _convert(value, encodings[index])
                            for index, value in enumerate(row)
                        ]
                    )
                    if len(current_rows) >= batch_size:
                        row_count += _write_batch(
                            writer,
                            arrow_schema,
                            current_rows,
                            expected_digest,
                            expected_nulls,
                            expected_empty,
                        )
                        current_rows.clear()
                        check_temporary_size()
                    if row_count + len(current_rows) > _positive_env(
                        "MAX_DATASET_ROWS", 1_000_000
                    ):
                        raise NormalizationError(
                            "resource_limit",
                            "The normalized output exceeds the configured dataset row limit.",
                        )
            finally:
                close_rows = getattr(table.rows, "close", None)
                if close_rows is not None:
                    close_rows()
            if current_rows:
                row_count += _write_batch(
                    writer,
                    arrow_schema,
                    current_rows,
                    expected_digest,
                    expected_nulls,
                    expected_empty,
                )
                check_temporary_size()
            writer.close()
            writer = None
        except NormalizationError:
            raise
        except (OSError, ValueError, TypeError, OverflowError, pa.ArrowException):
            raise NormalizationError(
                "conversion_failed",
                "Source values could not be represented by the normalized Parquet schema.",
            ) from None

        size = output_path.stat().st_size
        check_temporary_size()
        checksum = _file_sha256(output_path)
        validation = _validate_parquet(
            output_path,
            arrow_schema,
            row_count,
            expected_digest.hexdigest(),
            expected_nulls,
            expected_empty,
        )
        if detected_format == "parquet":
            with pq.ParquetFile(source_path) as source_file:
                if source_file.metadata.num_rows != validation["row_count"]:
                    raise NormalizationError(
                        "fidelity_validation_failed",
                        "Normalized Parquet row count differs from the input file.",
                    )

        object_key = (
            f"{dataset_id}/"
            f"{file_id}/"
            f"{source_checksum}/{config_hash}/"
            f"table-{table_index:04d}-{checksum}.parquet"
        )
        output = {
            "table_index": table_index,
            "table_name": table.name,
            "object_key": object_key,
            "file_size_bytes": size,
            "sha256": checksum,
            "compression": compression,
            "schema": validation["schema"],
            "row_count": validation["row_count"],
            "column_count": validation["column_count"],
            "null_value_counts": validation["null_value_counts"],
            "empty_value_counts": validation["empty_value_counts"],
            "rows_sha256": validation["rows_sha256"],
            "validation_status": "passed",
            "warnings": warnings,
        }

        if not storage.bucket_exists(processed_bucket):
            storage.make_bucket(processed_bucket)
        _put_verified_object(
            storage,
            processed_bucket,
            object_key,
            output_path,
            size,
            checksum,
        )
        return output, size
    finally:
        if writer is not None:
            writer.close()
        try:
            output_path.unlink()
        except FileNotFoundError:
            pass


def _write_batch(
    writer,
    schema: pa.Schema,
    rows: list[list],
    digest,
    null_counts: list[int],
    empty_counts: list[int],
) -> int:
    arrays = []
    for column_index, field in enumerate(schema):
        values = [row[column_index] for row in rows]
        try:
            arrays.append(pa.array(values, type=field.type, safe=True))
        except (pa.ArrowException, TypeError, ValueError, OverflowError):
            raise NormalizationError(
                "conversion_failed",
                f"Values in column {field.name!r} do not fit the safe Parquet type.",
            ) from None
    batch = pa.RecordBatch.from_arrays(arrays, schema=schema)
    for row_index in range(batch.num_rows):
        values = [
            batch.column(column_index)[row_index].as_py()
            for column_index in range(batch.num_columns)
        ]
        _row_digest_update(digest, values)
        for column_index, value in enumerate(values):
            if value is None:
                null_counts[column_index] += 1
            elif value == "":
                empty_counts[column_index] += 1
    writer.write_batch(batch)
    return batch.num_rows


def _file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        while chunk := source.read(1024 * 1024):
            digest.update(chunk)
    return digest.hexdigest()


def _put_verified_object(
    storage,
    bucket: str,
    key: str,
    path: Path,
    size: int,
    checksum: str,
):
    try:
        stat = storage.stat_object(bucket, key)
    except S3Error as error:
        if error.code not in {"NoSuchKey", "NoSuchObject"}:
            raise
    else:
        if stat.size == size and _remote_sha256(storage, bucket, key, size) == checksum:
            return

    with path.open("rb") as source:
        storage.put_object(
            bucket,
            key,
            source,
            size,
            content_type="application/vnd.apache.parquet",
            metadata={"sha256": checksum},
        )
    try:
        stat = storage.stat_object(bucket, key)
    except (MinioException, HTTPError, OSError):
        raise NormalizationError(
            "storage_verification_failed",
            "Normalized Parquet output could not be verified in object storage.",
        ) from None
    if stat.size != size or _remote_sha256(storage, bucket, key, size) != checksum:
        raise NormalizationError(
            "storage_verification_failed",
            "Normalized Parquet output size or checksum did not match after storage.",
        )


def _remote_sha256(storage, bucket: str, key: str, expected_size: int) -> str:
    response = storage.get_object(bucket, key)
    digest = hashlib.sha256()
    size = 0
    try:
        while chunk := response.read(1024 * 1024):
            size += len(chunk)
            if size > expected_size:
                return ""
            digest.update(chunk)
    finally:
        response.close()
        release = getattr(response, "release_conn", None)
        if release is not None:
            release()
    return digest.hexdigest() if size == expected_size else ""


def get_normalization_configuration() -> dict:
    compression, batch_size, max_output_bytes, config_hash = _normalization_settings()
    return {
        "compression": compression,
        "batch_size": batch_size,
        "max_output_bytes": max_output_bytes,
        "configuration_sha256": config_hash,
        "pipeline_version": NORMALIZATION_PIPELINE_VERSION,
    }


def verify_normalization_result(storage, bucket: str, result: dict) -> bool:
    for table in result.get("tables", []):
        key = table.get("object_key")
        if key is None:
            if table.get("row_count") != 0 or table.get("column_count") != 0:
                return False
            continue
        try:
            stat = storage.stat_object(bucket, key)
        except S3Error as error:
            if error.code in {"NoSuchKey", "NoSuchObject"}:
                return False
            raise
        size = table.get("file_size_bytes")
        checksum = table.get("sha256")
        if (
            not isinstance(size, int)
            or not isinstance(checksum, str)
            or stat.size != size
            or _remote_sha256(storage, bucket, key, size) != checksum
        ):
            return False
    return True


def validate_normalization_result(
    result: dict,
    parsed_result: dict,
    source_filename: str,
    source_size: int,
    source_checksum: str,
    dataset_id: int,
    file_id: int,
) -> dict:
    try:
        config = result["configuration"]
        source = result["source"]
        tables = result["tables"]
        if (
            not isinstance(config, dict)
            or not isinstance(source, dict)
            or result.get("version") != 1
            or isinstance(result.get("version"), bool)
            or result.get("pipeline_version") != NORMALIZATION_PIPELINE_VERSION
            or result.get("status") != "Ready"
            or result.get("detected_format") != parsed_result["detected_format"]
            or source
            != {
                "dataset_id": dataset_id,
                "file_id": file_id,
                "filename": source_filename,
                "sha256": source_checksum,
                "file_size_bytes": source_size,
            }
            or not isinstance(config.get("configuration_sha256"), str)
            or len(config["configuration_sha256"]) != 64
            or any(
                character not in "0123456789abcdef"
                for character in config["configuration_sha256"]
            )
            or not isinstance(config.get("compression"), str)
            or not isinstance(config.get("batch_size"), int)
            or isinstance(config.get("batch_size"), bool)
            or config["batch_size"] <= 0
            or config.get("pipeline_version") != NORMALIZATION_PIPELINE_VERSION
            or result.get("validation_status") != "passed"
            or not isinstance(result.get("row_count"), int)
            or isinstance(result.get("row_count"), bool)
            or not isinstance(result.get("column_count"), int)
            or isinstance(result.get("column_count"), bool)
            or not isinstance(result.get("output_size_bytes"), int)
            or isinstance(result.get("output_size_bytes"), bool)
            or not isinstance(result.get("warnings"), list)
            or any(not isinstance(value, str) for value in result["warnings"])
            or not isinstance(tables, list)
        ):
            raise ValueError("Normalization top-level metadata is inconsistent.")

        if parsed_result["detected_format"] == "xlsx":
            worksheets = parsed_result["metadata"].get("worksheets")
            if not isinstance(worksheets, list):
                raise ValueError("Parsed worksheet metadata is invalid.")
            expected_tables = [
                (
                    item["name"],
                    item["row_count"],
                    [column["name"] for column in item["columns"]],
                )
                for item in worksheets
            ]
        else:
            expected_columns = [
                column["name"] for column in parsed_result["columns"]
            ]
            if parsed_result["detected_format"] == "xml":
                container_columns, _container_values = _xml_container_fields(
                    parsed_result
                )
                expected_columns.extend(
                    column["name"] for column in container_columns
                )
            expected_tables = [
                (
                    None,
                    parsed_result["row_count"],
                    expected_columns,
                )
            ]
        if len(tables) != len(expected_tables):
            raise ValueError("Normalized table count does not match parsing metadata.")

        total_rows = 0
        total_columns = 0
        total_bytes = 0
        for index, (table, expected) in enumerate(zip(tables, expected_tables)):
            if not isinstance(table, dict):
                raise ValueError("Normalized table metadata is invalid.")
            name, rows, columns = expected
            schema = table["schema"]
            if not isinstance(schema, list) or any(
                not isinstance(field, dict)
                or not isinstance(field.get("name"), str)
                or not isinstance(field.get("type"), str)
                for field in schema
            ):
                raise ValueError("Normalized table schema is invalid.")
            if (
                table.get("table_index") != index
                or isinstance(table.get("table_index"), bool)
                or table.get("table_name") != name
                or table.get("row_count") != rows
                or isinstance(table.get("row_count"), bool)
                or table.get("column_count") != len(columns)
                or isinstance(table.get("column_count"), bool)
                or [field.get("name") for field in schema] != columns
                or table.get("validation_status") != "passed"
                or not isinstance(table.get("warnings"), list)
                or any(not isinstance(value, str) for value in table["warnings"])
            ):
                raise ValueError("Normalized table metadata is inconsistent.")
            null_counts = table["null_value_counts"]
            empty_counts = table["empty_value_counts"]
            if (
                not isinstance(null_counts, list)
                or len(null_counts) != len(columns)
                or any(
                    not isinstance(value, int)
                    or isinstance(value, bool)
                    or value < 0
                    for value in null_counts
                )
                or not isinstance(empty_counts, list)
                or len(empty_counts) != len(columns)
                or any(
                    not isinstance(value, int)
                    or isinstance(value, bool)
                    or value < 0
                    for value in empty_counts
                )
            ):
                raise ValueError("Normalized value counts are invalid.")
            key = table.get("object_key")
            if not columns and rows == 0:
                if (
                    key is not None
                    or table.get("file_size_bytes") != 0
                    or table.get("sha256") is not None
                ):
                    raise ValueError("Empty worksheet output is inconsistent.")
            else:
                checksum = table["sha256"]
                size = table["file_size_bytes"]
                expected_key = (
                    f"{dataset_id}/{file_id}/{source_checksum}/"
                    f"{config['configuration_sha256']}/"
                    f"table-{index:04d}-{checksum}.parquet"
                )
                if (
                    not isinstance(checksum, str)
                    or len(checksum) != 64
                    or any(character not in "0123456789abcdef" for character in checksum)
                    or not isinstance(size, int)
                    or isinstance(size, bool)
                    or size <= 0
                    or not isinstance(table.get("rows_sha256"), str)
                    or len(table["rows_sha256"]) != 64
                    or key != expected_key
                ):
                    raise ValueError("Normalized object metadata is invalid.")
                total_bytes += size
            total_rows += rows
            total_columns += len(columns)
        if (
            total_rows != parsed_result["row_count"]
            or total_rows != result.get("row_count")
            or total_columns != result.get("column_count")
            or total_bytes != result.get("output_size_bytes")
        ):
            raise ValueError("Normalized aggregate counts are inconsistent.")
        return json.loads(json.dumps(result, allow_nan=False))
    except (KeyError, TypeError, ValueError, OverflowError) as error:
        raise NormalizationError(
            "invalid_normalization_result",
            "Normalization returned incomplete or inconsistent result metadata.",
        ) from error


def normalize_file(
    source_path: str | Path,
    source_filename: str,
    parsed_result: dict,
    dataset_id: int,
    file_id: int,
    source_checksum: str,
    storage,
    processed_bucket: str,
) -> dict:
    compression, batch_size, max_output_bytes, config_hash = _normalization_settings()
    source_path = Path(source_path)

    try:
        with source_path.open("rb") as source:
            actual_parsed = parse_dataset_file(
                source_filename,
                source,
                source_path.stat().st_size,
            ).as_dict()
    except ParserConfigurationError:
        raise NormalizationError(
            "invalid_configuration",
            "Parser limits are invalid for normalization.",
        ) from None
    except (ParsingError, OSError):
        raise NormalizationError(
            "source_parse_failed",
            "The stored source could not be parsed again for normalization.",
        ) from None
    if actual_parsed != parsed_result:
        raise NormalizationError(
            "source_metadata_mismatch",
            "The stored source structure no longer matches its validated parsing metadata.",
        )

    tables = []
    output_bytes = 0
    table_sources = _table_sources(source_path, parsed_result)
    try:
        for table_index, table in enumerate(table_sources):
            if not table.columns and parsed_result["detected_format"] == "xlsx":
                tables.append(
                    {
                        "table_index": table_index,
                        "table_name": table.name,
                        "object_key": None,
                        "file_size_bytes": 0,
                        "sha256": None,
                        "compression": compression,
                        "schema": [],
                        "row_count": 0,
                        "column_count": 0,
                        "null_value_counts": [],
                        "empty_value_counts": [],
                        "validation_status": "passed",
                        "warnings": [
                            "Empty worksheet has no columns; its identity is retained in normalization metadata without a Parquet object."
                        ],
                    }
                )
                continue
            result, size = _normalize_table(
                table,
                source_path,
                parsed_result["detected_format"],
                dataset_id,
                file_id,
                source_checksum,
                config_hash,
                table_index,
                compression,
                batch_size,
                max_output_bytes - output_bytes,
                storage,
                processed_bucket,
            )
            output_bytes += size
            tables.append(result)
    finally:
        close_tables = getattr(table_sources, "close", None)
        if close_tables is not None:
            close_tables()

    if not tables:
        raise NormalizationError(
            "empty_source",
            "No logical tables were available to normalize.",
        )
    if sum(table["row_count"] for table in tables) != parsed_result["row_count"]:
        raise NormalizationError(
            "fidelity_validation_failed",
            "Normalized table row counts do not match the parsed source.",
        )

    return {
        "version": 1,
        "pipeline_version": NORMALIZATION_PIPELINE_VERSION,
        "status": "Ready",
        "source": {
            "dataset_id": dataset_id,
            "file_id": file_id,
            "filename": source_filename,
            "sha256": source_checksum,
            "file_size_bytes": source_path.stat().st_size,
        },
        "configuration": {
            "compression": compression,
            "batch_size": batch_size,
            "configuration_sha256": config_hash,
            "pipeline_version": NORMALIZATION_PIPELINE_VERSION,
        },
        "detected_format": parsed_result["detected_format"],
        "row_count": sum(table["row_count"] for table in tables),
        "column_count": sum(table["column_count"] for table in tables),
        "output_size_bytes": output_bytes,
        "tables": tables,
        "validation_status": "passed",
        "warnings": list(dict.fromkeys(
            warning for table in tables for warning in table["warnings"]
        )),
    }
